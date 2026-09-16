"""把经验证的数组回填为附件5要求的五份结果工作簿。"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "problems" / "选题C_2026正式" / "附件" / "附件5"
OUT_DIR = ROOT / "results" / "submission_results"
OUT_DIR.mkdir(parents=True, exist_ok=True)
TOL = 1e-7


def slot_labels() -> list[str]:
    wb = load_workbook(TEMPLATE_DIR / "result2.xlsx", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    labels = [str(ws.cell(1, c).value) for c in range(2, 146)]
    wb.close()
    if len(labels) != 144:
        raise ValueError("结果模板时段列不是144列")
    return labels


def style_sheet(ws, freeze="A2") -> None:
    blue = "1F4E78"
    light = "D9EAF7"
    thin = Side(style="thin", color="A7B8C5")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=blue))
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    ws.sheet_view.showGridLines = False


def new_sheet(wb, name, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    return ws


def add_readme(wb, model_text: str) -> None:
    ws = wb.create_sheet("填写说明", 0)
    rows = [
        ("项目", "说明"),
        ("单位", "购电量、充电量、放电量、储电量均为kWh；电费为元"),
        ("时段对齐", "原始附件与模板按同序号对应，不循环平移；模板末时段解释为次日0:00—0:10"),
        ("模型口径", model_text),
        ("数值精度", "工作簿保留底层浮点值，显示4位小数；连续紧急购电时段已合并"),
    ]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 95
    style_sheet(ws, "A2")


def aggregate_blocks(charge, discharge, soc):
    rows = []
    for start in range(0, 144, 24):
        block_index = start // 24
        state_time = "0:00" if block_index == 0 else ("24:00" if block_index == 1 else None)
        state_value = float(soc[0]) if block_index == 0 else (float(soc[-1]) if block_index == 1 else None)
        rows.append((
            f"{start // 6}:00-{(start + 24) // 6}:00",
            float(charge[start:start + 24].sum()),
            float(discharge[start:start + 24].sum()),
            state_time,
            state_value,
        ))
    return rows


def emergency_events(values, labels):
    events = []
    i = 0
    while i < 144:
        if values[i] <= TOL:
            i += 1
            continue
        j = i
        while j + 1 < 144 and values[j + 1] > TOL:
            j += 1
        start = labels[i].split("-")[0]
        end = labels[j].split("-")[-1]
        events.append((f"{start}-{end}", float(values[i:j + 1].sum())))
        i = j + 1
    return events


def export_q1() -> Path:
    data = np.genfromtxt(ROOT / "results" / "q1_detail.csv", delimiter=",", names=True, encoding="utf-8-sig")
    labels = slot_labels()
    wb = Workbook()
    wb.remove(wb.active)
    ws = new_sheet(wb, "计划购电量", ["时间段", "购电量"])
    for label, value in zip(labels, data["grid_kwh"]):
        ws.append([label, float(value)])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 18
    style_sheet(ws)
    ws = new_sheet(wb, "充放电量", ["时间段", "充电量", "放电量", "时刻", "储电量"])
    soc = np.r_[float(data["soc_start_kwh"][0]), data["soc_end_kwh"]]
    for row in aggregate_blocks(data["charge_kwh"], data["discharge_kwh"], soc):
        ws.append(row)
    for c, width in enumerate([20, 18, 18, 14, 18], 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    style_sheet(ws)
    path = OUT_DIR / "result1.xlsx"
    wb.save(path)
    return path


def export_annual(npz_path: Path, output_name: str, prefix: str = "", adjusted: bool = False) -> Path:
    z = np.load(npz_path, allow_pickle=True)
    def arr(name):
        return z[f"{prefix}{name}"]
    dates = np.asarray(z["dates"]).astype(str)
    start = int(np.flatnonzero(dates >= "2025-02-01")[0])
    labels = slot_labels()
    initial = arr("initial_grid") if f"{prefix}initial_grid" in z else arr("planned_grid_kwh")
    effective = arr("effective_grid") if f"{prefix}effective_grid" in z else initial
    charge = arr("charge") if f"{prefix}charge" in z else arr("planned_charge_kwh")
    discharge = arr("discharge") if f"{prefix}discharge" in z else arr("planned_discharge_kwh")
    emergency = arr("emergency") if f"{prefix}emergency" in z else arr("actual_emergency_kwh")
    if f"{prefix}soc" in z:
        soc = arr("soc")
    else:
        s0, s1 = arr("soc_start_kwh"), arr("soc_end_kwh")
        # Q2明细只存日首日末；按状态方程重建144个时段边界。
        soc = np.zeros((len(dates), 145))
        soc[:, 0] = s0
        soc[:, 1:] = s0[:, None] + np.cumsum(0.9 * charge - discharge / 0.9, axis=1)
        if np.max(np.abs(soc[:, -1] - s1)) > 1e-5:
            raise ValueError("Q2 SOC重建与日末值不一致")
    planned_cost = arr("planned_cost") if f"{prefix}planned_cost" in z else arr("planned_cost_yuan")
    emergency_cost = arr("emergency_cost") if f"{prefix}emergency_cost" in z else arr("emergency_cost_yuan")
    adjustment_cost = arr("adjustment_cost") if f"{prefix}adjustment_cost" in z else np.zeros(len(dates))
    total_cost = planned_cost + adjustment_cost + emergency_cost

    wb = Workbook()
    wb.remove(wb.active)
    model_text = "滚动调整MILP，计划、调整与紧急购电逐次结算" if adjusted else "严格日前两阶段MILP，实际阶段仅紧急购电与弃光"
    ws = new_sheet(wb, "计划购电量", ["日期\\时间", *labels, "全天购电量", "全天购电费"])
    for i in range(start, len(dates)):
        ws.append([dates[i], *map(float, initial[i]), float(initial[i].sum()), float(planned_cost[i])])
    ws.column_dimensions["A"].width = 14
    for c in range(2, 148):
        ws.column_dimensions[get_column_letter(c)].width = 13
    style_sheet(ws, "B2")

    if adjusted:
        ws = new_sheet(wb, "调整购电量", ["日期\\时间", *labels, "全天购电量", "全天购电费"])
        for i in range(start, len(dates)):
            ws.append([dates[i], *map(float, effective[i]), float(effective[i].sum()), float(total_cost[i])])
        ws.column_dimensions["A"].width = 14
        for c in range(2, 148):
            ws.column_dimensions[get_column_letter(c)].width = 13
        style_sheet(ws, "B2")

    ws = new_sheet(wb, "充放电量", ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for i in range(start, len(dates)):
        for block_index, block in enumerate(aggregate_blocks(charge[i], discharge[i], soc[i])):
            ws.append([dates[i] if block_index == 0 else None, *block])
    for c, width in enumerate([14, 20, 18, 18, 14, 18], 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    style_sheet(ws)

    ws = new_sheet(wb, "紧急购电量", ["日期", "购电时间段", "购电量"])
    for i in range(start, len(dates)):
        events = emergency_events(emergency[i], labels)
        for event_index, (interval, value) in enumerate(events):
            ws.append([dates[i] if event_index == 0 else None, interval, value])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 18
    style_sheet(ws)

    path = OUT_DIR / output_name
    wb.save(path)
    return path


def main() -> None:
    paths = [export_q1()]
    if (ROOT / "results" / "q2_detail.npz").exists():
        paths.append(export_annual(ROOT / "results" / "q2_detail.npz", "result2.xlsx"))
    if (ROOT / "results" / "q3_detail.npz").exists():
        paths.append(export_annual(ROOT / "results" / "q3_detail.npz", "result3.xlsx", adjusted=True))
    if (ROOT / "results" / "q4_detail.npz").exists():
        paths.append(export_annual(ROOT / "results" / "q4_detail.npz", "result4-2.xlsx", prefix="q42_"))
        paths.append(export_annual(ROOT / "results" / "q4_detail.npz", "result4-3.xlsx", prefix="q43_", adjusted=True))
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
