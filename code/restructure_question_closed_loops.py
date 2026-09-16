"""将C题论文重构为逐问闭环的国奖论文式章节顺序。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_国奖行文强化版.docx"
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_分问闭环重构版.docx"


def find(doc: Document, prefix: str) -> Paragraph:
    hits = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        raise ValueError(f"定位失败: {prefix!r}, hits={len(hits)}")
    return hits[0]


def replace_number_label(paragraph: Paragraph, old: str, new: str) -> None:
    """只改公式右侧编号文本，不触碰同段中的Office数学公式对象。"""
    changed = False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            changed = True
    if not changed:
        raise ValueError(f"公式编号未找到: {old} in {paragraph.text!r}")


def first_rpr(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def replace_text(paragraph: Paragraph, text: str, *, bold: bool | None = None) -> None:
    rpr = first_rpr(paragraph)
    ppr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not ppr:
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    if bold is not None:
        run.bold = bold


def clone_paragraph(template: Paragraph, text: str, style: str | None = None, *, bold: bool = False) -> Paragraph:
    elem = OxmlElement("w:p")
    if template._p.pPr is not None:
        elem.append(deepcopy(template._p.pPr))
    paragraph = Paragraph(elem, template._parent)
    if style:
        paragraph.style = style
    rpr = first_rpr(template)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    run.bold = bold
    return paragraph


def insert_after(anchor: Paragraph, new_paragraph: Paragraph) -> Paragraph:
    anchor._p.addnext(new_paragraph._p)
    return new_paragraph


def insert_before(anchor: Paragraph, new_paragraph: Paragraph) -> Paragraph:
    anchor._p.addprevious(new_paragraph._p)
    return new_paragraph


def remove_between(start: Paragraph, end: Paragraph) -> None:
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def nodes_between(start: Paragraph, end: Paragraph) -> list:
    nodes = []
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nodes.append(node)
        node = node.getnext()
    return nodes


def move_nodes_after(nodes: list, anchor: Paragraph) -> Paragraph:
    cursor = anchor._p
    for node in nodes:
        cursor.addnext(node)
        cursor = node
    return Paragraph(cursor, anchor._parent)


def find_image_paragraph(doc: Document, target_ref: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        for rel_id in paragraph._p.xpath(".//a:blip/@r:embed"):
            if doc.part.rels[rel_id].target_ref == target_ref:
                return paragraph
    raise ValueError(f"找不到图片: {target_ref}")


def replace_in_runs(doc: Document, mapping: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            original = run.text
            if not original:
                # 图片、域代码等非文本运行不能重写，否则会丢失其XML子节点。
                continue
            value = original
            for old, new in mapping.items():
                value = value.replace(old, new)
            if value != original:
                run.text = value


def force_heading_bold(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            for run in paragraph.runs:
                run.bold = True


def anonymize(doc: Document) -> None:
    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category", "identifier"):
        setattr(props, attr, "")


def main() -> None:
    doc = Document(INPUT)
    normal = find(doc, "本文针对光伏出力")
    h1_template = find(doc, "一、问题重述")
    h2_template = find(doc, "1.1 问题背景")

    # 删除集中式“问题分析/模型选型”章节；分析回到各问题内部。
    old_analysis = find(doc, "二、问题分析")
    old_assumptions = find(doc, "三、模型假设与数据处理")
    remove_between(old_analysis, old_assumptions)
    old_analysis._p.getparent().remove(old_analysis._p)

    # 调整公共章节顺序：模型假设 → 符号说明 → 数据预处理。
    replace_text(old_assumptions, "二、模型假设", bold=True)
    replace_text(find(doc, "3.1基本假设"), "2.1 基本假设", bold=True)
    time_heading = find(doc, "3.2时间与单位对齐")
    symbol_heading = find(doc, "四、符号说明")
    common_heading = find(doc, "五、由实际运行过程建立基本调度关系")
    time_nodes = [time_heading._p] + nodes_between(time_heading, symbol_heading)
    replace_text(symbol_heading, "三、符号说明", bold=True)
    data_h1 = clone_paragraph(h1_template, "四、数据预处理与探索性分析", "Heading 1", bold=True)
    common_heading._p.addprevious(data_h1._p)
    data_h1 = Paragraph(data_h1._p, old_assumptions._parent)
    move_nodes_after(time_nodes, data_h1)
    replace_text(time_heading, "4.1 时间与单位对齐", bold=True)
    last_time = Paragraph(time_nodes[-1], old_assumptions._parent)
    data_h2 = insert_after(last_time, clone_paragraph(h2_template, "4.2 数据质量与基本特征", "Heading 2", bold=True))
    p = insert_after(data_h2, clone_paragraph(normal,
        "对附件1—4逐表检查日期、时段、缺失值和取值范围。附件1、附件2及附件4均无缺失值和负值；附件3共有1095个日期空白，均来自同一日期四个预报时刻只在首行填写日期的版式设计，故在组内向下填充，而不作为观测缺失插补。原始功率观测经范围核查后不作主观删改，统一乘以1/6小时转换为每个10分钟时段的电量；附件3整点预测插值到10分钟网格后仅作非负裁剪。"))
    insert_after(p, clone_paragraph(normal,
        "探索性统计表明，附件1的负荷为3309.39—5958.97 kW，光伏最高7612.32 kW，净负荷有30个时段为负，说明部分日间时段存在可用于充电的光伏富余；分时电价为0.3713—1.3952元/kWh，峰谷差使储能具有跨时段转移价值。附件2覆盖365个连续日期，附件1负荷和光伏曲线分别与附件2逐时刻全年均值高度一致，因此问题一可视为典型日调度，问题二至问题四则按自然日顺序滚动计算。"))

    # 把公共物理关系并入问题一的“模型建立”，不再单列统一模型。
    q1_h1 = find(doc, "六、问题一")
    common_nodes = nodes_between(common_heading, q1_h1)
    common_intro = find(doc, "本章只写四问都会反复使用")
    common_nodes = [n for n in common_nodes if n is not common_intro._p]
    common_intro._p.getparent().remove(common_intro._p)
    common_heading._p.getparent().remove(common_heading._p)
    # 保存后续需要移动并按新章节顺序重排的公式编号段落。
    eq5 = find(doc, "(5)")
    eq6 = find(doc, "(6)")
    eq7 = find(doc, "(7)")
    eq10 = find(doc, "(10)")
    replace_text(q1_h1, "五、问题一的模型建立与求解", bold=True)
    q1_analysis_h = insert_after(q1_h1, clone_paragraph(h2_template, "5.1 问题分析", "Heading 2", bold=True))
    q1_analysis = insert_after(q1_analysis_h, clone_paragraph(normal,
        "问题一给出某日全部144个时段的电价、居民负荷和光伏出力，要求在满足供电与储能安全边界的条件下使全天购电费用最小。光伏可以直接供负荷，也可以给储能充电；储能又能把当前电量转移到后续高价或光伏不足时段。因此，本问不能按单个时段独立分配电源，而要先写出母线电量平衡和储电量递推，再将全天时段联立求解。"))
    q1_model_h = insert_after(q1_analysis, clone_paragraph(h2_template, "5.2 模型建立", "Heading 2", bold=True))
    move_nodes_after(common_nodes, q1_model_h)
    for prefix, text in [
        ("5.1 每个时段", "5.2.1 时段电量平衡"),
        ("5.2 储电量", "5.2.2 储能状态、边界与充放电互斥"),
        ("5.3 日末", "5.2.3 初始储电量条件"),
    ]:
        h = find(doc, prefix)
        h.style = "Heading 3"
        replace_text(h, text, bold=True)
    q1_objective_h = find(doc, "6.1 由全天已知条件")
    q1_objective_h.style = "Heading 3"
    replace_text(q1_objective_h, "5.2.4 目标函数与完整模型", bold=True)
    q1_results_h = find(doc, "6.2 最优计划")
    replace_text(q1_results_h, "5.3 模型求解与计算结果", bold=True)
    q1_analysis_result = find(doc, "图6-1把结果中的三条逻辑")
    insert_before(q1_analysis_result, clone_paragraph(h2_template, "5.4 结果分析与检验", "Heading 2", bold=True))
    replace_text(q1_analysis_result,
        "图5-1显示，低价或光伏富余时储电量上升，电价较高且负荷超过光伏时储能放电并压低购电量；储电量触及1200 kWh或10800 kWh边界后，后续动作受容量限制。计划购电量的突变与分时电价台阶和5000 kW功率上限一致。逐时复算表明供需平衡和储电量递推残差均低于10⁻¹² kWh，且不存在同时充放电；连续松弛与混合整数模型同值，说明本组数据的最优解自然满足互斥。")

    # 问题二：分析 → 数据/场景 → 模型 → 结果 → 分析检验。
    q2_h1 = find(doc, "七、问题二")
    replace_text(q2_h1, "六、问题二的模型建立与求解", bold=True)
    q2_h_analysis = find(doc, "7.1 为什么必须区分")
    replace_text(q2_h_analysis, "6.1 问题分析", bold=True)
    q2_intro = find(doc, "附件2给出了全年的实际负荷")
    replace_text(q2_intro,
        "问题二仍要求最小化正常购电费与紧急购电费之和，但每天0:00制定计划时，当天真实负荷和光伏尚未发生。此时可用信息只有当前储电量、过去日期的实际记录和题定电价；计划冻结后，题目只允许在供电不足时按5倍价格紧急购电，不能重新安排储能。因此，本问必须先用历史数据描述当天可能出现的净负荷，再把决策分为0时计划和事后结算两个阶段。")
    info_nodes = [find(doc, "在同一个决策日的所有可能实际情景")._p]
    info_nodes.insert(0, find(doc, "(8)")._p)
    # 公式段落的可见文字只有编号，定位后连同两阶段说明与式(9)移入模型建立。
    p82 = find(doc, "在同一个决策日的所有可能实际情景")
    p84 = find(doc, "第二阶段不再允许重新安排储能")
    node = p82._p
    moved_q2 = []
    # 加入位于p82之前的信息集公式段。
    previous = node.getprevious()
    if previous is not None:
        moved_q2.append(previous)
    while node is not None:
        moved_q2.append(node)
        if node is p84._p:
            break
        node = node.getnext()
    q2_data_h = find(doc, "7.2 只用过去数据")
    replace_text(q2_data_h, "6.2 数据处理与情景构造", bold=True)
    q2_model_h = find(doc, "7.3 计划时权衡")
    replace_text(q2_model_h, "6.3 两阶段模型建立", bold=True)
    move_nodes_after(moved_q2, q2_model_h)
    # 跨日状态传递属于全年顺序运行的Q2—Q4，不放在Q1模型中。
    cross_start = find(doc, "Q2—Q4按自然时间顺序运行")
    cross_end = find(doc, "问题一明确要求单日首尾")
    cross_nodes = []
    node = cross_start._p
    while node is not None:
        cross_nodes.append(node)
        if node is cross_end._p:
            break
        node = node.getnext()
    q2_stage_end = find(doc, "第二阶段不再允许重新安排储能")
    move_nodes_after(cross_nodes, q2_stage_end)
    replace_text(cross_end,
        "问题二至问题四均按自然日期连续运行，不设置人为的每日复位。上一日末储电量原样传给下一日；1月从6000 kWh起始值顺序运行，以得到2月1日的真实初始状态。每日0:00求计划时采用48小时延伸窗估计日末电量的后续价值，但只执行前24小时，避免在当天末端无代价放空储能。")
    q2_result_h = find(doc, "7.4 求解结果")
    replace_text(q2_result_h, "6.4 模型求解与计算结果", bold=True)
    q2_result = find(doc, "2月1日至12月31日，计划购电费用")
    replace_text(q2_result,
        "2月1日至12月31日，计划购电费用为13750414.64元，紧急购电费用为1507608.25元，总费用为15258022.89元；紧急购电量为404258.73 kWh，预测加权绝对百分比误差为10.95%。表6-1—表6-3给出指定日期的计划购电、储能动作和紧急购电，图6-1与图6-2分别展示代表日的预测偏差及月度费用分解。")
    q3_h1 = find(doc, "八、问题三")
    q2_eval_h = insert_before(q3_h1, clone_paragraph(h2_template, "6.5 结果分析与检验", "Heading 2", bold=True))
    insert_after(q2_eval_h, clone_paragraph(normal,
        "结果显示，正常购电覆盖了绝大部分需求，紧急费用约占总费用的9.88%；紧急购电主要出现在实际净负荷高于日前场景覆盖范围的时段，说明剩余费用来自信息不足而非供电约束失效。按实际曲线逐时结算后，母线平衡、储电量边界和跨日状态传递均满足要求；训练数据截止日期也均早于决策日，未使用当天未来实际值。"))

    # 问题三：分析 → 预测处理 → 滚动模型 → 结果 → 分析检验。
    replace_text(q3_h1, "七、问题三的模型建立与求解", bold=True)
    q3_analysis_h = find(doc, "8.1 新预报")
    replace_text(q3_analysis_h, "7.1 问题分析", bold=True)
    q3_analysis = find(doc, "问题三与问题二的根本差别")
    replace_text(q3_analysis,
        "问题三在问题二基础上，于0:00、6:00、12:00和18:00提供新的未来光伏预报。新信息到达时，过去时段已经执行、当前储电量已经形成，二者均不能改写；只有尚未执行的购电和储能计划可以调整。因此，本问需要以每次预报发布时刻为决策节点，重复完成“读取当前状态—更新未来输入—优化剩余时域—执行至下一节点”，并按题目规定逐次结算增购和减购费用。")
    q3_preprocess_para = find(doc, "每次求解先把附件3的整点光伏预测")
    insert_before(q3_preprocess_para, clone_paragraph(h2_template, "7.2 预测数据处理与状态更新", "Heading 2", bold=True))
    q3_model_h = find(doc, "8.2非对称调整费用")
    replace_text(q3_model_h, "7.3 滚动优化模型建立", bold=True)
    insert_after(q3_model_h, clone_paragraph(normal,
        "每次滚动求解均继承问题一的能量平衡、储能状态和互斥约束，并以当前实测储电量作为新初值。设调整前仍有效的购电计划为原计划，新优化得到的未执行时段购电量为新计划；两者的正差计为增购，负差的绝对值计为减购，由此得到非对称调整费用。"))
    q3_result_h = find(doc, "8.3 是否值得更新")
    replace_text(q3_result_h, "7.4 模型求解与计算结果", bold=True)
    insert_after(q3_result_h, clone_paragraph(normal,
        "全年计算得到：仅在0:00制定计划的总费用为16391391.76元，采用0:00、6:00、12:00和18:00四次更新后的总费用为15731687.04元，调整净费用为499519.56元，紧急购电量为748796.98 kWh。图7-1、图7-2和表7-2—表7-4给出更新频率、代表日预测修正及指定日期的购电和储能结果。"))
    q3_analysis_result = find(doc, "与仅在0:00制定计划相比")
    insert_before(q3_analysis_result, clone_paragraph(h2_template, "7.5 结果分析与检验", "Heading 2", bold=True))
    replace_text(q3_analysis_result,
        "与仅0:00计划相比，四次更新使总费用减少659704.73元，降幅约4.02%，紧急购电量下降28.03%。这说明新预报减少的高价应急损失超过了逐次调整成本，故四次更新在本组数据上具有实际收益。逐版本复算表明，每次增购、减购均相对上一版有效计划结算，已执行时段保持冻结，母线平衡、储电量边界和状态递推均满足约束。")

    # 问题四：分析 → 价格处理 → 模型 → 结果 → 信息价值分析。
    q4_h1 = find(doc, "九、问题四")
    replace_text(q4_h1, "八、问题四的模型建立与求解", bold=True)
    q4_analysis_h = insert_after(q4_h1, clone_paragraph(h2_template, "8.1 问题分析", "Heading 2", bold=True))
    q4_analysis = insert_after(q4_analysis_h, clone_paragraph(normal,
        "问题四把问题二、三中已知的分时电价改为实时价格。附件4给出的全年价格是事后实现路径，不能视为每天0:00或日内更新时已经知道的未来信息。因此，本问首先要在每个决策时刻利用历史价格预测未实现时段，再分别沿用问题二的日前两阶段结构和问题三的滚动结构；同时设置事后信息基准，用于衡量价格未知造成的费用损失。"))
    q4_price_h = find(doc, "9.1 实时价格不能直接")
    replace_text(q4_price_h, "8.2 价格数据处理与因果预测", bold=True)
    q4_model_h = find(doc, "9.2 区分现实可执行")
    replace_text(q4_model_h, "8.3 调度模型与评价基准", bold=True)
    q4_model_intro = find(doc, "因果策略给出在当时信息条件")
    replace_text(q4_model_intro,
        "将每个决策时刻得到的价格预测替换原模型中的已知电价：问题4-2每天0:00预测全天价格并求解两阶段计划，问题4-3在四次光伏预报更新时同步更新价格，只调整尚未执行时段。二者均按当天实际价格进行事后费用结算。为识别价格预测误差和总体信息不足的影响，另设置以下两类评价基准：")
    q4_picture = find_image_paragraph(doc, "media/image7.png")
    q4_result_h = insert_before(q4_picture, clone_paragraph(h2_template, "8.4 模型求解与计算结果", "Heading 2", bold=True))
    insert_after(q4_result_h, clone_paragraph(normal,
        "因果价格预测下，问题4-2与问题4-3的全年总费用分别为16124373.46元和16550967.30元；对应的完全价格信息反事实费用分别为15923572.41元和16401924.52元，全外生信息严格下界均为12780453.36元。图8-1和表8-2汇总不同信息条件下的费用及指定日期结果。"))
    q4_analysis_result = find(doc, "因果策略相对完全价格信息基准")
    insert_before(q4_analysis_result, clone_paragraph(h2_template, "8.5 结果分析与信息价值", "Heading 2", bold=True))
    replace_text(q4_analysis_result,
        "问题4-2和问题4-3相对完全价格信息基准分别增加200801.05元和149042.78元，反映价格未知带来的样本内代价。问题4-3的因果总费用比问题4-2高426593.84元，表明本样本中更频繁的调整虽然获得更新信息，但增减购电成本和剩余紧急购电仍超过其收益。严格下界仅用于刻画同时放宽未来价格、负荷、光伏信息和整数互斥后的理论边界，不作为现实可执行方案。")

    # 终章统一编号，并补充推广小节。
    verification_h = find(doc, "十、模型检验")
    replace_text(verification_h, "九、模型检验与结果可靠性", bold=True)
    replace_text(find(doc, "10.1硬约束"), "9.1 硬约束与费用账本检验", bold=True)
    replace_text(find(doc, "10.2时间信息"), "9.2 时间信息泄漏检验", bold=True)
    replace_text(find(doc, "10.3模型对照"), "9.3 模型对照检验", bold=True)
    evaluation_h = find(doc, "十一、模型评价")
    replace_text(evaluation_h, "十、模型评价、改进与推广", bold=True)
    replace_text(find(doc, "11.1优点"), "10.1 模型优点", bold=True)
    replace_text(find(doc, "11.2局限"), "10.2 模型局限与改进", bold=True)
    ai_h = find(doc, "十二、AI工具使用声明")
    promotion_h = insert_before(ai_h, clone_paragraph(h2_template, "10.3 模型推广", "Heading 2", bold=True))
    insert_after(promotion_h, clone_paragraph(normal,
        "本文框架可推广到含风电、可控负荷和多类储能的园区微网。推广时只需在母线平衡中增加相应能量流，在状态方程中加入设备特性，并依据实际市场规则改写购售电与偏差结算项；“已执行决策不可回写、预测只使用当时可得数据”的信息边界仍保持不变。若获得电池衰减和需量电费参数，还可在目标函数中加入寿命成本与最大需量费用。"))
    replace_text(ai_h, "十一、AI工具使用声明", bold=True)
    replace_text(find(doc, "十三、参考文献"), "十二、参考文献", bold=True)
    replace_text(find(doc, "十四、附录"), "十三、附录 完整源程序代码", bold=True)

    # 章节移动后同步更新图表编号和正文引用。
    replace_in_runs(doc, {
        "图6-": "图5-", "表6-": "表5-",
        "图7-": "图6-", "表7-": "表6-",
        "图8-": "图7-", "表8-": "表7-",
        "图9-": "图8-", "表9-": "表8-",
    })
    # 新增结果段已使用重构后的目标编号，需在旧编号映射后恢复正确引用。
    replace_text(q2_result,
        "2月1日至12月31日，计划购电费用为13750414.64元，紧急购电费用为1507608.25元，总费用为15258022.89元；紧急购电量为404258.73 kWh，预测加权绝对百分比误差为10.95%。表6-1—表6-3给出指定日期的计划购电、储能动作和紧急购电，图6-1与图6-2分别展示代表日的预测偏差及月度费用分解。")
    replace_text(find(doc, "全年计算得到"),
        "全年计算得到：仅在0:00制定计划的总费用为16391391.76元，采用0:00、6:00、12:00和18:00四次更新后的总费用为15731687.04元，调整净费用为499519.56元，紧急购电量为748796.98 kWh。图7-1、图7-2和表7-2—表7-4给出更新频率、代表日预测修正及指定日期的购电和储能结果。")
    replace_text(find(doc, "因果价格预测下"),
        "因果价格预测下，问题4-2与问题4-3的全年总费用分别为16124373.46元和16550967.30元；对应的完全价格信息反事实费用分别为15923572.41元和16401924.52元，全外生信息严格下界均为12780453.36元。图8-1和表8-2汇总不同信息条件下的费用及指定日期结果。")
    # 公式按正文出现次序连续编号：Q1为(1)—(6)，Q2从数据处理的(7)开始。
    replace_number_label(eq6, "(6)", "(5)")
    replace_number_label(eq7, "(7)", "(6)")
    replace_number_label(eq10, "(10)", "(7)")
    replace_number_label(eq5, "(5)", "(10)")
    force_heading_bold(doc)
    anonymize(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
