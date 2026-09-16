from __future__ import annotations

import json
from pathlib import Path

from docx import Document


DOCX = Path(r"D:\数模工作流\delivery\C题论文_光伏微网储能购电滚动优化_全文格式与参考文献终修版.docx")
OUT = Path(r"D:\数模工作流\state\agent_outputs\c_paper_visual_audit.json")


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


doc = Document(DOCX)
paras = doc.paragraphs
figures = []
for i, p in enumerate(paras):
    text = p.text.strip()
    if has_drawing(p) or text.startswith("图"):
        lo = max(0, i - 2)
        hi = min(len(paras), i + 4)
        figures.append(
            {
                "paragraph_index": i,
                "has_drawing": has_drawing(p),
                "text": text,
                "context": [
                    {
                        "index": j,
                        "style": paras[j].style.name if paras[j].style else "",
                        "text": paras[j].text.strip(),
                        "has_drawing": has_drawing(paras[j]),
                    }
                    for j in range(lo, hi)
                ],
            }
        )

tables = []
for ti, table in enumerate(doc.tables):
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    tables.append(
        {
            "table_index": ti,
            "rows": len(rows),
            "cols": max((len(r) for r in rows), default=0),
            "preview": rows[:4],
        }
    )

payload = {
    "source": str(DOCX),
    "paragraph_count": len(paras),
    "inline_shape_count": len(doc.inline_shapes),
    "table_count": len(doc.tables),
    "figure_and_caption_contexts": figures,
    "tables": tables,
}
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(OUT)
print(f"paragraphs={len(paras)} figures/captions={len(figures)} inline_shapes={len(doc.inline_shapes)} tables={len(doc.tables)}")
