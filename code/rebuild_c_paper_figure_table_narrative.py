from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\数模工作流")
SOURCE = Path(r"E:\微信聊天记录\xwechat_files\wxid_so1zh5t7c8rl22_8e76\msg\file\2026-09\C题论文_光伏微网储能购电滚动优化_全文格式与参考文献终修版(1)(1).docx")
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_图表逐项解读版.docx"
FIG = ROOT / "results" / "figures_matlab"
AUDIT = ROOT / "state" / "agent_outputs" / "c_paper_figure_table_narrative_audit.json"
ACTIVE_DOC = None

FIGURES = {
    "图6-1典型日功率、储能动作与SOC轨迹": ("图6-1 典型日功率、储能动作与SOC轨迹", FIG / "q1_dispatch_matlab_awardstyle.png"),
    "图7-1四个代表日的日前预测、实际净负荷与紧急购电": ("图7-1 四个代表日的日前预测、实际净负荷与紧急购电", FIG / "q2_representative_days_matlab.png"),
    "图7-2正式期月度费用分解与紧急购电量": ("图7-2 正式期月度费用分解与紧急购电量", FIG / "q2_monthly_cost_emergency_matlab.png"),
    "图8-1不同日内更新频率的费用与紧急购电量对比": ("图8-1 不同日内更新频率的费用与紧急购电量对比", FIG / "q3_update_ablation_matlab.png"),
    "图8-2四个指定日的滚动预测修正效果": ("图8-2 四个指定日的滚动预测修正效果", FIG / "q3_forecast_updates_matlab.png"),
    "图9-1因果策略、完全价格信息基准与全信息下界的费用对比": ("图9-1 因果策略、完全价格信息基准与全信息下界的费用对比", FIG / "q4_information_value_matlab.png"),
}

TABLE_CAPTIONS = [
    "表 6-1 问题一全天调度指标与对照", "表6-2指定时段购电量及全天汇总", "表6-3储能分段充/放电量及首末储电量",
    "表7-1指定日期、指定时段的购电量及全天汇总", "表7-2指定日期储能分段充/放电量（单元格为充电量/放电量，kWh）", "表7-3指定日期紧急购电",
    "表 8-1 不同预报更新频率的全年代价对照", "表8-2指定日期、指定时段的0时计划量/最终有效量及全天汇总", "表8-3指定日期储能分段充/放电量（单元格为充电量/放电量，kWh）", "表8-4指定日期紧急购电",
    "表 9-1 波动电价下三类策略的费用对照", "表9-2波动电价下指定日期结果",
]


def iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def exact_paragraph(doc, text: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"找不到段落：{text}")


def starting_paragraph(doc, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"找不到以此开头的段落：{prefix}")


def has_drawing(p: Paragraph) -> bool:
    return bool(p._p.xpath(".//w:drawing | .//w:pict"))


def neighboring_element(doc, caption: Paragraph, kind: str, prefer_previous: bool = False):
    blocks = list(iter_blocks(doc))
    idx = next(i for i, b in enumerate(blocks) if b._element is caption._p)
    for radius in range(1, 5):
        order = (idx - radius, idx + radius) if prefer_previous else (idx + radius, idx - radius)
        for j in order:
            if 0 <= j < len(blocks):
                b = blocks[j]
                if kind == "drawing" and isinstance(b, Paragraph) and has_drawing(b):
                    return b
                if kind == "table" and isinstance(b, Table):
                    return b
    raise ValueError(f"{caption.text}附近找不到{kind}")


def font_run(run, size: float = 12.0, bold: bool | None = None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:eastAsia"), "宋体")


def body_format(p: Paragraph, keep_next: bool = False):
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.keep_with_next = keep_next
    pf.widow_control = True
    for r in p.runs:
        font_run(r, 12)


def caption_format(p: Paragraph, table: bool):
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3 if table else 2)
    pf.space_after = Pt(2 if table else 6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = table
    for r in p.runs:
        font_run(r, 10.5)


def figure_format(p: Paragraph):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = True


def set_text(p: Paragraph, text: str, mode: str = "body"):
    p.clear()
    p.add_run(text)
    if mode == "body":
        body_format(p)
    else:
        caption_format(p, table=(mode == "table_caption"))


def new_paragraph(text: str, mode: str = "body") -> Paragraph:
    el = OxmlElement("w:p")
    p = Paragraph(el, ACTIVE_DOC._body)
    p.add_run(text)
    if mode == "body":
        body_format(p)
    else:
        caption_format(p, table=(mode == "table_caption"))
    return p


def new_figure(path: Path, width_cm: float = 15.6) -> Paragraph:
    if not path.exists():
        raise FileNotFoundError(path)
    p = new_paragraph("")
    p.clear()
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    figure_format(p)
    return p


def replace_picture(p: Paragraph, path: Path, width_cm: float = 15.6):
    if not path.exists():
        raise FileNotFoundError(path)
    p.clear()
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    figure_format(p)


def insert_sequence(anchor_el, objects):
    cur = anchor_el
    for obj in objects:
        el = obj._element if hasattr(obj, "_element") else obj
        cur.addnext(el)
        cur = el
    return cur


def clear_between(start_el, stop_el):
    node = start_el.getnext()
    while node is not None and node is not stop_el:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def set_cell_border(cell, edge: str, val: str, size: str = "8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    tag = "w:" + edge
    element = borders.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        borders.append(element)
    element.set(qn("w:val"), val)
    if val != "nil":
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), "000000")


def format_table(table: Table):
    cols = len(table.columns)
    presets = {
        2: [9.8, 5.8], 3: [2.7, 10.2, 2.7], 4: [3.9] * 4, 5: [3.1, 3.1, 3.1, 3.1, 3.2],
        6: [2.5, 2.7, 2.5, 2.7, 2.5, 2.7], 8: [2.0, 1.8, 1.8, 1.8, 1.8, 1.8, 1.8, 2.8],
        9: [2.0, 1.5, 1.5, 1.5, 1.5, 1.5, 1.5, 2.3, 2.3],
    }
    widths = presets.get(cols, [15.6 / cols] * cols)
    size = 8.0 if cols >= 9 else 8.5 if cols >= 8 else 9.0 if cols >= 6 else 9.5
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ri, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if ri == 0 and tr_pr.find(qn("w:tblHeader")) is None:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            tr_pr.append(hdr)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[min(ci, len(widths) - 1)])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.left_indent = pf.right_indent = pf.first_line_indent = Pt(0)
                pf.space_before = pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in p.runs:
                    font_run(run, size, bold=(ri == 0))
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                set_cell_border(cell, edge, "nil")
            if ri == 0:
                set_cell_border(cell, "top", "single", "12")
                set_cell_border(cell, "bottom", "single", "8")
            if ri == len(table.rows) - 1:
                set_cell_border(cell, "bottom", "single", "12")


def normalize_caption_text(text: str) -> str:
    text = re.sub(r"^(图|表)\s+(\d+-\d+)", r"\1\2", text.strip())
    return re.sub(r"^((?:图|表)\d+-\d+)\s*", r"\1 ", text)


def main():
    global ACTIVE_DOC
    doc = Document(SOURCE)
    ACTIVE_DOC = doc
    source_sha = __import__("hashlib").sha256(SOURCE.read_bytes()).hexdigest()

    # Capture all original table and figure objects before rebuilding the result blocks.
    table_objs = {}
    for old in TABLE_CAPTIONS:
        cap = exact_paragraph(doc, old)
        table_objs[normalize_caption_text(old)] = (
            cap,
            neighboring_element(doc, cap, "table", prefer_previous=old.startswith("表 8-1")),
        )

    figure_objs = {}
    for old, (new, path) in FIGURES.items():
        cap = exact_paragraph(doc, old)
        holder = neighboring_element(
            doc,
            cap,
            "drawing",
            prefer_previous=old.startswith("图8-") or old.startswith("图9-"),
        )
        set_text(cap, new, "figure_caption")
        replace_picture(holder, path)
        figure_objs[new] = (cap, holder)

    # Normalize table caption text while preserving its paragraph objects.
    for new, (cap, _table) in table_objs.items():
        set_text(cap, new, "table_caption")

    q1_h = exact_paragraph(doc, "6.2 模型求解与计算结果")
    q1_anchor = starting_paragraph(doc, "将144个时段的购电、充电、放电、弃光、储电量及充放电状态按时间顺序组成决策向量")
    q1_stop = exact_paragraph(doc, "6.3 结果分析与检验")
    clear_between(q1_anchor._p, q1_stop._p)
    q1 = []
    c, t = table_objs["表6-1 问题一全天调度指标与对照"]
    q1 += [new_paragraph("为先判断储能是否真正降低全天费用，并核对整数互斥约束是否改变最优值，表6-1将主模型、连续松弛和无储能基线放在同一口径下比较。"), c, t,
           new_paragraph("表6-1表明，主模型日购电费用为35126.95元，较无储能方案的48052.05元减少12925.10元，降幅为26.90%。连续松弛与混合整数模型的费用完全相同，说明这组数据下连续最优解已经自然满足充放电互斥，整数变量主要承担物理可行性的显式保证。")]
    c, t = table_objs["表6-2 指定时段购电量及全天汇总"]
    q1 += [new_paragraph("在确认总体经济性后，还需检查模型如何在题目指定的关键时刻分配购电量。表6-2列出六个10分钟时段及全天汇总。"), c, t,
           new_paragraph("表6-2中10:00、14:00和20:00的购电量为0，而18:00达到636.98 kWh。这些差异由光伏出力、电价台阶、储能荷电状态和5000 kW功率上限共同决定，并非人为规定能源使用先后顺序。")]
    c, t = table_objs["表6-3 储能分段充/放电量及首末储电量"]
    q1 += [new_paragraph("为解释上述购电差异如何由跨时段储能实现，表6-3进一步汇总各四小时时段的充放电量，并核对首末储电量。"), c, t,
           new_paragraph("表6-3显示，储能在0:00—4:00和光伏较充足的8:00—16:00累计吸收电量，在4:00—8:00及16:00—20:00集中释放；24:00储电量回到6000.00 kWh，与0:00相同，因而节省费用不是通过透支日末库存取得。")]
    c, f = figure_objs["图6-1 典型日功率、储能动作与SOC轨迹"]
    q1 += [new_paragraph("表格给出精确数值后，图6-1把供需功率、储能动作和荷电状态放到同一时间轴上，以检验调度动作是否与光伏和电价变化相吻合。"), f, c,
           new_paragraph("图6-1(a)显示光伏能够覆盖负荷时外网购电明显降低；图6-1(b)中充电集中在低价或光伏富余时段，放电集中在高价且净负荷较高时段；图6-1(c)表明荷电状态始终位于10%—90%边界内，并在日末回到初值。三幅子图共同形成“供需变化—储能响应—状态可行”的证据闭环。")]
    insert_sequence(q1_anchor._p, q1)

    q2_anchor = starting_paragraph(doc, "2月1日至12月31日，计划购电费用为13750414.64元")
    set_text(q2_anchor, "2月1日至12月31日，计划购电费用为13750414.64元，紧急购电费用为1507608.25元，总费用为15258022.89元；紧急购电量为404258.73 kWh，预测加权绝对百分比误差为10.95%。下文依次从指定时点结果、代表日偏差、月度费用结构和跨日状态连续性四个层面展示计算结果。")
    q2_stop = exact_paragraph(doc, "7.4 结果分析与检验")
    clear_between(q2_anchor._p, q2_stop._p)
    q2 = []
    c, t = table_objs["表7-1 指定日期、指定时段的购电量及全天汇总"]
    q2 += [new_paragraph("首先用表7-1回答题目指定日期和时段的日前购电安排，并给出各代表日的全天购电量与费用。"), c, t,
           new_paragraph("表7-1显示，四个代表日的购电规模随季节净负荷变化明显：光伏较强的6月21日全天购电量最低，而3月20日和12月21日更高。各时点数值均是0:00冻结计划的组成部分，不含事后紧急购电。")]
    c, t = table_objs["表7-2 指定日期储能分段充/放电量（单元格为充电量/放电量，kWh）"]
    q2 += [new_paragraph("由于计划购电量必须与储能动作共同解释，表7-2按六个四小时时段汇总充放电，并列出每日首末储电量。"), c, t,
           new_paragraph("表7-2表明，不同日期的储能动作会随预测净负荷改变，但每天24:00的状态均由当天真实执行结果决定并传递到次日。首末值并不被强制设为相等，从而避免把全年问题错误地拆成彼此独立的单日优化。")]
    c, t = table_objs["表7-3 指定日期紧急购电"]
    q2 += [new_paragraph("计划冻结后，只有实际净负荷超过既定供给的时段才产生紧急购电。表7-3列出指定日期中出现缺口的时间区间及其电量。"), c, t,
           new_paragraph("表7-3中的紧急购电集中于若干连续短时段，说明高额费用来自局部预测偏差，而不是全天供电能力不足。该表也验证了第二阶段只对未覆盖缺口结算，没有在实际值揭示后重新安排储能。")]
    c, f = figure_objs["图7-1 四个代表日的日前预测、实际净负荷与紧急购电"]
    q2 += [new_paragraph("为直观看出紧急购电为何发生，图7-1叠加四个代表日的日前预测、实际净负荷和紧急购电位置。"), f, c,
           new_paragraph("图7-1显示，当实际净负荷在午后或傍晚高于日前预测，而既定购电与储能放电又无法覆盖偏差时，红色紧急购电随即出现。6月21日偏差较小，全天紧急购电仅196.52 kWh；其余代表日的尖峰与净负荷低估时段基本对应，符合严格两阶段结算逻辑。")]
    c, f = figure_objs["图7-2 正式期月度费用分解与紧急购电量"]
    q2 += [new_paragraph("代表日只能解释局部机理，因此图7-2把正式期各月的计划费用、紧急费用和紧急购电量统一汇总，以观察季节差异。"), f, c,
           new_paragraph("图7-2表明计划购电费始终构成月度费用主体，紧急购电费约占全年总费用的9.88%。6月和7月的紧急购电量相对较高，说明光伏波动与净负荷预测误差的季节性会显著影响事后缺口，但并未改变正常购电承担主体供能的总体结构。")]
    q2_soc = new_figure(FIG / "q2_soc_continuity_matlab.png")
    q2_soc_cap = new_paragraph("图7-3 正式期跨日SOC连续性及日内运行范围", "figure_caption")
    q2 += [new_paragraph("最后，费用下降只有在储能状态全年连续且不越界时才具有可信度。图7-3绘制每天的日初、日末及日内最小和最大储电量。"), q2_soc, q2_soc_cap,
           new_paragraph("图7-3显示，正式期最低和最高储电量分别为1997.08 kWh和8550.00 kWh，均处于1200—10800 kWh允许区间；每天24:00与下一天0:00逐点相接，最大连接残差为0 kWh。因此全年费用不是通过每日重置或跨日状态断裂获得。")]
    insert_sequence(q2_anchor._p, q2)

    q3_anchor = starting_paragraph(doc, "全年计算得到：仅在0:00制定计划的总费用为16391391.76元")
    set_text(q3_anchor, "全年计算得到：仅在0:00制定计划的总费用为16391391.76元，采用0:00、6:00、12:00和18:00四次更新后的总费用为15731687.04元，调整净费用为499519.56元，紧急购电量为748796.98 kWh。下文先用消融对照判断更新是否值得，再展示预测修正过程和指定日期的最终执行结果。")
    q3_stop = exact_paragraph(doc, "8.4 结果分析与检验")
    clear_between(q3_anchor._p, q3_stop._p)
    q3 = []
    c, t = table_objs["表8-1 不同预报更新频率的全年代价对照"]
    q3 += [new_paragraph("是否采用全部更新时刻不能凭直觉判断，必须比较新增信息带来的应急费用下降与调整成本。表8-1依次加入6:00、12:00和18:00更新，形成可比的年度消融实验。"), c, t,
           new_paragraph("表8-1显示，更新时刻由仅0:00增加到四次后，总费用从16391391.76元降至15731687.04元，节省659704.73元；紧急购电量从1040382.80 kWh降至748796.98 kWh，下降28.03%。尽管调整费用上升到499519.56元，但紧急购电损失的减少更大，因此完整更新在本组数据上仍有净收益。")]
    c, f = figure_objs["图8-1 不同日内更新频率的费用与紧急购电量对比"]
    q3 += [new_paragraph("为同时观察总费用、调整费用和紧急购电量的变化方向，图8-1将表8-1的核心指标转化为同一组视觉对照。"), f, c,
           new_paragraph("图8-1直观显示，总费用与紧急购电量随更新次数增加而单调下降，但边际收益逐步减小：加入18:00更新后总费用仅再下降21786.47元，而调整费用继续增加。由此可见，第四次更新虽仍有效，却已接近收益递减区间。")]
    c, f = figure_objs["图8-2 四个指定日的滚动预测修正效果"]
    q3 += [new_paragraph("年度消融说明更新有效，但还需检验它如何改变日内决策。图8-2展示四个指定日期中0:00、6:00、12:00和18:00预测对剩余净负荷轨迹的逐步修正。"), f, c,
           new_paragraph("图8-2中，每条更新曲线只从其发布时间向后延伸，发布时间之前的已执行区间不被改写；后续预测逐步贴近实际轨迹，使剩余购电计划能及时修正。该图同时证明滚动策略遵守因果信息边界，而不是利用整日实际值事后重算。")]
    c, t = table_objs["表8-2 指定日期、指定时段的0时计划量/最终有效量及全天汇总"]
    q3 += [new_paragraph("预测修正最终会落实为购电计划变化。表8-2将指定时点的0时计划量与最后一次更新后的有效量并列，便于核对调整方向和幅度。"), c, t,
           new_paragraph("表8-2显示，不同时段的最终有效购电量既可能上调也可能下调，且变化集中在更新时刻之后。结果说明模型没有把“更新”机械理解为增加购电，而是依据新预测重新平衡正常购电、储能和潜在紧急缺口。")]
    c, t = table_objs["表8-3 指定日期储能分段充/放电量（单元格为充电量/放电量，kWh）"]
    q3 += [new_paragraph("购电计划变化必须与真实可执行的储能轨迹一致，因此表8-3给出相同日期的分段充放电量和首末荷电状态。"), c, t,
           new_paragraph("表8-3表明，各日储能在新预报到达后重新分配剩余时段的充放电，但始终从当时实测状态继续递推。首末状态差异反映跨日能量转移，而非约束遗漏；逐时检查确认容量、功率和互斥条件均满足。")]
    c, t = table_objs["表8-4 指定日期紧急购电"]
    q3 += [new_paragraph("为检验更新后仍未消除的风险，表8-4列出指定日期最终需要紧急购电的区间与电量。"), c, t,
           new_paragraph("表8-4说明滚动更新能够削减但不能完全消除紧急购电。剩余缺口主要出现在最后一次可用预报之后或实际净负荷突变较大的时段，符合“信息改善而非获得完全信息”的题意。")]
    insert_sequence(q3_anchor._p, q3)

    q4_h = exact_paragraph(doc, "9.3 模型求解与计算结果")
    q4_stop = exact_paragraph(doc, "9.4 结果分析与信息价值")
    clear_between(q4_h._p, q4_stop._p)
    q4 = [new_paragraph("因果价格预测下，问题4-2与问题4-3的全年总费用分别为16124373.46元和16550967.30元。为区分现实可执行结果、价格预测损失和全部未来信息价值，下文依次比较年度费用层级，并给出指定日期的实际结算结果。")]
    c, t = table_objs["表9-1 波动电价下三类策略的费用对照"]
    q4 += [new_paragraph("表9-1把因果在线策略、只预知未来价格的反事实基准和同时预知价格、负荷、光伏的严格下界置于同一费用口径下。"), c, t,
           new_paragraph("表9-1中，问题4-2和问题4-3的因果总费用分别比完全价格信息基准高200801.05元和149042.78元，对应1.261%和0.909%。全外生信息下界为12780453.36元，明显更低，但它同时使用现实决策时不可获得的未来价格、负荷和光伏，仅能作为理论边界。")]
    c, f = figure_objs["图9-1 因果策略、完全价格信息基准与全信息下界的费用对比"]
    q4 += [new_paragraph("由于价格信息损失相对全年总费用较小，单看表格不易辨认其量级。图9-1先展示三类信息条件下的成本层级，再单独放大因价格未知造成的差额。"), f, c,
           new_paragraph("图9-1(a)显示两种因果策略均高于完全价格信息基准和全信息下界；图9-1(b)进一步表明价格预测损失只占完全价格信息费用约1%。这意味着储能和滚动计划已经缓冲了大部分价格不确定性，剩余较大的成本差距主要来自负荷、光伏等其他未来信息不可知。")]
    c, t = table_objs["表9-2 波动电价下指定日期结果"]
    q4 += [new_paragraph("年度汇总之外，表9-2列出四个指定日期在问题4-2与问题4-3下的总费用和紧急购电量，以检查结论是否由个别日期异常主导。"), c, t,
           new_paragraph("表9-2显示，滚动价格与光伏更新在不同日期上的收益方向并不一致：部分日期降低费用，另一些日期因调整结算和预测偏差反而增加费用。由此可知，问题4-3全年费用高于问题4-2并非模型失效，而是题定非对称调整成本在波动电价下累积后的结果。")]
    insert_sequence(q4_h._p, q4)

    # Rebuild each result-analysis section as genuine synthesis rather than delayed figure captions.
    synthesis = {
        "6.3 结果分析与检验": (
            "七、问题二的模型建立与求解",
            [
                "综合上述结果，问题一的最优策略不是预先规定光伏、储能和外网的固定使用顺序，而是在逐时供需平衡与跨时段储电量递推的共同约束下，由费用最小化自动形成削峰填谷。储能把低价或光伏富余时段的电量转移到高价缺电时段，使日费用下降26.90%，同时保持日末库存不变，因此节省具有可持续性。",
                "从可行性看，逐时母线平衡和储电量递推的最大残差均低于10⁻¹² kWh，SOC从未越过1200—10800 kWh边界，且不存在同时充放电。连续松弛与混合整数模型同值进一步表明，互斥约束没有通过牺牲目标值来换取物理可行性。故问题一的费用结果、运行机理和约束校验相互一致。",
            ],
        ),
        "7.4 结果分析与检验": (
            "八、问题三的模型建立与求解",
            [
                "问题二的全年结果说明，严格日前计划能够承担绝大部分供能任务，但预测误差会在计划冻结后转化为高价紧急购电。计划购电费为13750414.64元，紧急购电费为1507608.25元，后者占总费用9.88%；404258.73 kWh的紧急购电量集中在实际净负荷超出日前覆盖范围的局部时段，符合题目限定的唯一事后补救方式。",
                "模型按自然日期连续运行，日末SOC直接传给次日，全年连接残差为0 kWh；同时，所有预测样本、残差场景和参数选择均截止于决策日前，没有使用当天未来实际值。结合母线平衡、容量边界和预测误差检验，可确认费用差异来自真实的信息不足，而非每日重置储能或未来数据泄漏。",
            ],
        ),
        "8.4 结果分析与检验": (
            "九、问题四的模型建立与求解",
            [
                "问题三的关键结论是，新预报是否值得使用必须由“应急损失减少量”与“调整结算成本”共同决定。四次更新使总费用减少659704.73元、紧急购电量下降28.03%，表明在本组数据中信息收益超过调整成本；但新增18:00更新的边际节省已明显缩小，说明更新频率继续提高未必经济。",
                "逐次账本检查表明，每次增购和减购均相对上一版仍有效计划结算，已执行时段保持冻结，最终费用由计划购电、调整净费用和紧急购电共同构成。所有滚动节点均从当时实测SOC出发，状态递推、容量功率边界和充放电互斥均满足要求。因此该策略既利用了日内新信息，又保持了真实可执行性。",
            ],
        ),
        "9.4 结果分析与信息价值": (
            "十、模型检验与结果可靠性",
            [
                "问题四表明，附件4中的实时价格必须作为事后实现值而不是事前已知量处理。在这一因果口径下，问题4-2和问题4-3的价格预测损失仅为各自完全价格信息基准的1.261%和0.909%，说明历史价格预测已经捕捉主要波动结构，储能也对剩余误差形成缓冲。",
                "另一方面，问题4-3全年总费用比问题4-2高426593.84元，说明更频繁的滚动调整并不必然更优。波动电价下，1.5倍增购结算与按原计划价格退回50%的减购规则会放大交易成本，新增信息带来的缺口削减不足以覆盖该成本。完全价格信息基准用于单独衡量价格预测代价，全外生信息下界用于刻画理论最优边界，二者均不冒充现实可执行方案。",
            ],
        ),
    }
    for heading_text, (next_heading_text, paras) in synthesis.items():
        h = exact_paragraph(doc, heading_text)
        stop = exact_paragraph(doc, next_heading_text)
        clear_between(h._p, stop._p)
        insert_sequence(h._p, [new_paragraph(x) for x in paras])

    for table in doc.tables:
        format_table(table)

    # Preserve the source's full abstract verbatim; clean only metadata.
    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category", "identifier"):
        try:
            setattr(props, attr, "")
        except Exception:
            pass

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    AUDIT.write_text(json.dumps({
        "source": str(SOURCE), "source_sha256": source_sha, "output": str(OUTPUT),
        "source_abstract_preserved": True, "matlab_figures_replaced": 6, "matlab_figures_added": 1,
        "tables_reordered_with_local_explanations": 12, "result_synthesis_sections_rewritten": 4,
        "inline_shapes": len(doc.inline_shapes), "tables": len(doc.tables),
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)
    print(AUDIT)
    print(f"inline_shapes={len(doc.inline_shapes)} tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
