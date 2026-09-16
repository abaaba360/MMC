"""Synchronize the final paper appendix with the lean support package."""

from __future__ import annotations

import io
import keyword
import re
import shutil
import tokenize
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_图表逐项解读版.docx"
OUTPUT = ROOT / "delivery" / "C题论文.docx"
SUPPORT = ROOT / "state" / "submission_staging" / "support"
BACKUP = ROOT / "archive" / "提交前备份_20260912" / SOURCE.name


CODE_FILES = [
    ("Q1.py", "问题一确定性调度、连续松弛与无储能对照"),
    ("Q2.py", "问题二严格日前两阶段随机调度"),
    ("Q3.py", "问题三光伏预报更新下的滚动调度"),
    ("Q4.py", "问题四因果实时价格策略与信息基准"),
    ("common.py", "四问共用的能量平衡与储能调度求解器"),
    ("data_q1.py", "问题一附件读取与单位转换"),
    ("data.py", "全年数据读取、字段校验与单位转换"),
    ("forecast.py", "因果负荷、光伏、价格预测及残差场景生成"),
    ("rolling.py", "问题三、四的滚动状态更新与费用结算函数"),
    ("check.py", "四问能量平衡、状态转移、互斥和费用账本复核"),
    ("export.py", "向附件5结果工作簿回填计算结果"),
    ("export_matlab.py", "把已验证结果导出为MATLAB绘图数据"),
    ("run_all.py", "统一核验与完整复算入口"),
    ("plot_q1.m", "问题一调度与储能状态图的MATLAB程序"),
    ("plot_figures.m", "问题二至问题四证据图的MATLAB程序"),
]


def set_font(run, name: str, size: float, color: str = "000000", bold: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def set_cell_margins(cell, top=25, start=55, bottom=25, end=55) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_table_borders(table, color="B7B7B7", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def replace_scenario_claim(doc: Document) -> None:
    for p in doc.paragraphs:
        if p.text.startswith("问题二中，情景数量按残差分布统计量"):
            p.text = (
                "问题二中，情景数量不是预先指定的固定常数，而是每天根据决策日前可得的历史残差样本量，"
                "在10、20、30个候选值中自动选取。正式期334个决策日中，331日采用30个情景，2日采用10个情景，"
                "1日采用20个情景。该结果表明，在历史样本充分时算法保留更多整日残差轨迹；样本较少时则自动缩减"
                "情景数，避免在有限历史上人为扩充情景。因而本文只把可由逐日选择日志直接复核的情景数分布作为"
                "稳健性证据，不再给出缺少独立计算记录的固定情景数费用差异。"
            )
            for r in p.runs:
                set_font(r, "宋体", 10.5)
            return
    raise RuntimeError("未找到10.4情景数量段落")


def remove_after_paragraph(doc: Document, marker: str) -> None:
    paragraph = next((p for p in doc.paragraphs if p.text.strip() == marker), None)
    if paragraph is None:
        raise RuntimeError(f"未找到段落：{marker}")
    body = doc._element.body
    children = list(body)
    start = children.index(paragraph._p) + 1
    for child in children[start:]:
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def add_body_paragraph(doc: Document, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, "宋体", 10.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, "宋体", 10.5)
    else:
        r = p.add_run(text)
        set_font(r, "宋体", 10.5)
    return p


def add_support_table(doc: Document) -> None:
    rows = [
        ("Q1.py、Q2.py、Q3.py、Q4.py", "四个问题的主程序"),
        ("common.py、data_q1.py、data.py", "公共求解器及数据读取模块"),
        ("forecast.py、rolling.py", "因果预测、情景生成及滚动调度模块"),
        ("check.py", "四问结果独立核验程序"),
        ("export.py、export_matlab.py", "结果工作簿与MATLAB绘图数据导出程序"),
        ("plot_q1.m、plot_figures.m", "论文结果图的MATLAB绘图程序"),
        ("run_all.py、requirements.txt", "统一运行入口及Python依赖清单"),
        ("results/", "五份结果工作簿、四问JSON及必要CSV、NPZ中间结果"),
        ("figures/", "论文采用的8幅中文命名图片"),
        ("代码运行说明.txt", "快速核验、完整复算和绘图步骤"),
        ("AI工具使用详情.pdf", "AI工具名称、用途、交互、采纳修改与人工核验记录"),
        ("文件清单.json", "压缩包文件名与字节数清单"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(7.0)
    table.columns[1].width = Cm(8.0)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("文件名", "功能说明")):
        set_cell_fill(cell, "D9E2F3")
        set_cell_margins(cell, 75, 100, 75, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), "宋体", 10.5, bold=True)
    for left, right in rows:
        cells = table.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, (left, right))):
            set_cell_margins(cell, 55, 90, 55, 90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(text), "宋体", 9.5)
    set_table_borders(table)
    doc.add_paragraph()


PY_COLORS = {
    tokenize.COMMENT: "008000",
    tokenize.STRING: "008080",
    tokenize.NUMBER: "7030A0",
    tokenize.OP: "000000",
}


def python_colored_spans(text: str) -> dict[int, list[tuple[int, int, str]]]:
    spans: dict[int, list[tuple[int, int, str]]] = {}
    try:
        tokens = tokenize.generate_tokens(io.StringIO(text).readline)
        for tok in tokens:
            if tok.type in {tokenize.ENCODING, tokenize.ENDMARKER, tokenize.NEWLINE, tokenize.NL, tokenize.INDENT, tokenize.DEDENT}:
                continue
            color = PY_COLORS.get(tok.type)
            if tok.type == tokenize.NAME and keyword.iskeyword(tok.string):
                color = "0000CC"
            elif tok.type == tokenize.NAME and tok.string in {"True", "False", "None"}:
                color = "0000CC"
            if color and tok.start[0] == tok.end[0]:
                spans.setdefault(tok.start[0], []).append((tok.start[1], tok.end[1], color))
    except (tokenize.TokenError, IndentationError):
        pass
    return spans


def add_colored_line(paragraph, line: str, spans: list[tuple[int, int, str]], language: str) -> None:
    if language == "matlab":
        comment = line.find("%")
        if comment >= 0:
            spans = [(comment, len(line), "008000")]
        for match in re.finditer(r"\b(if|else|elseif|end|for|while|function|return|clear|close|clc)\b", line):
            spans.append((match.start(), match.end(), "0000CC"))
    spans = sorted(spans, key=lambda x: (x[0], -(x[1] - x[0])))
    pos = 0
    for start, end, color in spans:
        if start < pos or start > len(line):
            continue
        if start > pos:
            set_font(paragraph.add_run(line[pos:start]), "Consolas", 7.5)
        set_font(paragraph.add_run(line[start:end]), "Consolas", 7.5, color=color)
        pos = min(end, len(line))
    if pos < len(line):
        set_font(paragraph.add_run(line[pos:]), "Consolas", 7.5)
    if not line:
        set_font(paragraph.add_run(" "), "Consolas", 7.5)


def add_code_file(doc: Document, appendix_index: int, filename: str, description: str) -> None:
    heading = doc.add_paragraph()
    heading.style = doc.styles["Heading 2"]
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    set_font(heading.add_run(f"A.{appendix_index} {filename}"), "黑体", 12, bold=True)
    desc = doc.add_paragraph()
    desc.paragraph_format.keep_with_next = True
    desc.paragraph_format.space_before = desc.paragraph_format.space_after = Pt(0)
    lang = "MATLAB" if filename.endswith(".m") else "Python"
    set_font(desc.add_run(f"语言：{lang}；作用：{description}。以下代码与支撑材料中的同名文件逐字一致。"), "宋体", 9)

    text = (SUPPORT / filename).read_text(encoding="utf-8")
    lines = text.splitlines()
    spans_by_line = python_colored_spans(text) if filename.endswith(".py") else {}
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.9)
    table.columns[1].width = Cm(14.1)
    set_table_borders(table, color="C8C8C8", size="3")
    for number, line in enumerate(lines, 1):
        cells = table.add_row().cells
        cells[0].width = Cm(0.9)
        cells[1].width = Cm(14.1)
        set_cell_fill(cells[0], "F2F2F2")
        for cell in cells:
            set_cell_margins(cell, 0, 45, 0, 45)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p_num = cells[0].paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_num.paragraph_format.space_before = p_num.paragraph_format.space_after = Pt(0)
        p_num.paragraph_format.line_spacing = 1.0
        set_font(p_num.add_run(str(number)), "Consolas", 7, color="808080")
        p_code = cells[1].paragraphs[0]
        p_code.paragraph_format.space_before = p_code.paragraph_format.space_after = Pt(0)
        p_code.paragraph_format.line_spacing = 1.0
        add_colored_line(p_code, line, spans_by_line.get(number, []), "matlab" if filename.endswith(".m") else "python")
    doc.add_paragraph()


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not SUPPORT.exists():
        raise FileNotFoundError(SUPPORT)
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    doc = Document(SOURCE)
    replace_scenario_claim(doc)
    remove_after_paragraph(doc, "A.1 附件与支撑材料文件列表")

    add_body_paragraph(
        doc,
        "支撑材料按“主程序—必要公共模块—核验与导出—结果与图件—使用说明”的顺序组织，"
        "不包含赛题原始附件、论文旧稿、排版脚本或本机路径。文件清单如下。",
    )
    add_support_table(doc)
    add_body_paragraph(
        doc,
        "运行方式：先执行“python run_all.py --check”核验既有结果；如需完整复算，将赛题附件文件夹放在支撑材料根目录，"
        "再执行“python run_all.py --full”。MATLAB图由plot_q1.m与plot_figures.m生成。详细环境和步骤见代码运行说明.txt。",
        bold_prefix="运行方式：",
    )

    for idx, (filename, description) in enumerate(CODE_FILES, start=2):
        add_code_file(doc, idx, filename, description)

    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.keywords = ""
    core.subject = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
