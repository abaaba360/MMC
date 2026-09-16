from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"D:\数模工作流")
INPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_全文格式与参考文献终修版.docx"
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_图表论证增强版.docx"
FIG = ROOT / "results" / "figures_matlab"
AUDIT = ROOT / "state" / "agent_outputs" / "c_paper_visual_evidence_edit_audit.json"

REPLACEMENTS = {
    "图6-1典型日功率、储能动作与SOC轨迹": (
        "图6-1 典型日功率、储能动作与SOC轨迹",
        FIG / "q1_dispatch_matlab_awardstyle.png",
    ),
    "图7-1四个代表日的日前预测、实际净负荷与紧急购电": (
        "图7-1 四个代表日的日前预测、实际净负荷与紧急购电",
        FIG / "q2_representative_days_matlab.png",
    ),
    "图7-2正式期月度费用分解与紧急购电量": (
        "图7-2 正式期月度费用分解与紧急购电量",
        FIG / "q2_monthly_cost_emergency_matlab.png",
    ),
    "图8-1不同日内更新频率的费用与紧急购电量对比": (
        "图8-1 不同日内更新频率的费用与紧急购电量对比",
        FIG / "q3_update_ablation_matlab.png",
    ),
    "图8-2四个指定日的滚动预测修正效果": (
        "图8-2 四个指定日的滚动预测修正效果",
        FIG / "q3_forecast_updates_matlab.png",
    ),
    "图9-1因果策略、完全价格信息基准与全信息下界的费用对比": (
        "图9-1 因果策略、完全价格信息基准与全信息下界的费用对比",
        FIG / "q4_information_value_matlab.png",
    ),
}


def set_font(run, size: float, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


def normalize_body(paragraph) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.widow_control = True
    for run in paragraph.runs:
        set_font(run, 12)


def normalize_caption(paragraph) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = True
    for run in paragraph.runs:
        set_font(run, 10.5)


def normalize_figure_holder(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = True


def set_text(paragraph, text: str, caption: bool = False) -> None:
    paragraph.clear()
    paragraph.add_run(text)
    (normalize_caption if caption else normalize_body)(paragraph)


def insert_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    normalize_body(new_para)
    return new_para


def find_exact(doc: Document, text: str):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"找不到段落：{text}")


def preceding_drawing(doc: Document, caption) -> object:
    idx = next(i for i, p in enumerate(doc.paragraphs) if p._p is caption._p)
    for p in reversed(doc.paragraphs[max(0, idx - 3):idx]):
        if p._p.xpath(".//w:drawing"):
            return p
    raise ValueError(f"题注前未找到图片：{caption.text}")


def replace_picture(paragraph, path: Path, width_cm: float = 15.6) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph.clear()
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    normalize_figure_holder(paragraph)


def set_cell_border(cell, edge: str, val: str, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    tag = "w:" + edge
    element = borders.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        borders.append(element)
    element.set(qn("w:val"), val)
    if val != "nil":
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), "000000")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def column_widths(cols: int) -> list[float]:
    presets = {
        2: [9.7, 5.9],
        3: [2.7, 10.3, 2.6],
        4: [3.9, 3.9, 3.9, 3.9],
        5: [3.1, 3.1, 3.1, 3.1, 3.2],
        6: [2.5, 2.7, 2.5, 2.7, 2.5, 2.7],
        8: [2.1, 1.75, 1.75, 1.75, 1.75, 1.75, 1.75, 2.7],
        9: [2.0, 1.5, 1.5, 1.5, 1.5, 1.5, 1.5, 2.3, 2.3],
    }
    return presets.get(cols, [15.6 / cols] * cols)


def format_table(table) -> None:
    cols = len(table.columns)
    widths = column_widths(cols)
    size = 8.0 if cols >= 9 else 8.5 if cols >= 8 else 9.0 if cols >= 6 else 9.5
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ri, row in enumerate(table.rows):
        set_cant_split(row)
        if ri == 0:
            set_repeat_header(row)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[min(ci, len(widths) - 1)])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.left_indent = Pt(0)
                pf.right_indent = Pt(0)
                pf.first_line_indent = Pt(0)
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in p.runs:
                    set_font(run, size, bold=(ri == 0))
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                set_cell_border(cell, edge, "nil")
            if ri == 0:
                set_cell_border(cell, "top", "single", "12")
                set_cell_border(cell, "bottom", "single", "8")
            if ri == len(table.rows) - 1:
                set_cell_border(cell, "bottom", "single", "12")


def main() -> None:
    doc = Document(INPUT)
    replaced = []
    for old_caption, (new_caption, path) in REPLACEMENTS.items():
        caption = find_exact(doc, old_caption)
        holder = preceding_drawing(doc, caption)
        replace_picture(holder, path)
        set_text(caption, new_caption, caption=True)
        replaced.append({"caption": new_caption, "image": str(path.relative_to(ROOT))})

    # 问题一：图前明确观察目标，图后将各面板证据与表中数值分开解释。
    p = find_exact(doc, "将144个时段的购电、充电、放电、弃光、储电量及充放电状态按时间顺序组成决策向量，把式（1）—（6）写成稀疏线性约束，采用HiGHS分支定界算法求解混合整数线性规划。求解完成后按原始时间顺序还原各能量流，并逐时复算平衡、状态和互斥约束；同时求解无储能基线与连续松弛模型，用于评价储能收益和整数约束的实际影响。")
    p.add_run(" 为判断储能如何改变购电时序，并同时检查SOC边界，图6-1把供需功率、充放电动作及SOC—电价关系按同一时间轴排列。")
    normalize_body(p)
    p = find_exact(doc, "图6-1显示，低价或光伏富余时储电量上升，电价较高且负荷超过光伏时储能放电并压低购电量；储电量触及1200 kWh或10800 kWh边界后，后续动作受容量限制。计划购电量的突变与分时电价台阶和5000 kW功率上限一致。逐时复算表明供需平衡和储电量递推残差均低于10⁻¹² kWh，且不存在同时充放电；连续松弛与混合整数模型同值，说明本组数据的最优解自然满足互斥。")
    set_text(p, "图6-1(a)表明，光伏能够覆盖负荷时外网购电显著降低；图6-1(b)把每个10分钟时段的放电画在零轴上方、充电画在零轴下方，可见储能主要在低价或光伏富余时吸收电量，并在高价且净负荷较高时释放。图6-1(c)进一步显示SOC始终处于10%—90%的运行区间，充放电方向与电价变化及容量边界相互吻合。")
    insert_after(p, "表6-1和表6-2给出指定时段与全天精确数值。MILP日购电费用为35126.95元，较无储能基线减少12925.10元，即降低26.90%；LP连续松弛与MILP同值，逐时复算的供需平衡和储电量递推残差均低于10⁻¹² kWh，且不存在同时充放电，说明本组数据的最优连续解自然满足互斥。")

    # 问题二：补充全年SOC连续性图，并把三类证据逐图解释。
    p = find_exact(doc, "2月1日至12月31日，计划购电费用为13750414.64元，紧急购电费用为1507608.25元，总费用为15258022.89元；紧急购电量为404258.73 kWh，预测加权绝对百分比误差为10.95%。表7-1—表7-3给出指定日期的计划购电、储能动作和紧急购电，图7-1与图7-2分别展示代表日的预测偏差及月度费用分解。")
    set_text(p, "2月1日至12月31日，计划购电费用为13750414.64元，紧急购电费用为1507608.25元，总费用为15258022.89元；紧急购电量为404258.73 kWh，预测加权绝对百分比误差为10.95%。表7-1—表7-3回答指定日期和时段的精确结果；图7-1用于定位预测偏差与紧急购电，图7-2用于比较月度费用结构，图7-3用于检查全年SOC可行性与跨日连续性。")
    heading = find_exact(doc, "7.4 结果分析与检验")
    caption = heading.insert_paragraph_before("图7-3 正式期跨日SOC连续性及日内运行范围")
    normalize_caption(caption)
    holder = caption.insert_paragraph_before()
    replace_picture(holder, FIG / "q2_soc_continuity_matlab.png")

    p = find_exact(doc, "结果显示，正常购电覆盖了绝大部分需求，紧急费用约占总费用的9.88%；紧急购电主要出现在实际净负荷高于日前场景覆盖范围的时段，说明剩余费用来自信息不足而非供电约束失效。按实际曲线逐时结算后，母线平衡、储电量边界和跨日状态传递均满足要求；训练数据截止日期也均早于决策日，未使用当天未来实际值。")
    set_text(p, "图7-1显示，代表日中紧急购电集中于实际净负荷高于日前点预测且既定购电、充放电计划无法完全覆盖的时段；这正是第二阶段只允许以紧急购电补足缺口的结果，而不是重新调度储能。负净荷时段表示光伏出力超过居民负荷，不应解释为供电不足。")
    p2 = insert_after(p, "图7-2表明，计划购电费构成全年费用主体，紧急购电费占总费用9.88%；紧急购电量在6月和7月较高，说明同一预测方法在不同季节的缺口风险并不均匀。图7-3给出的日内SOC范围始终位于10%—90%的运行边界内，正式期最低和最高储电量分别为1997.08 kWh和8550.00 kWh。")
    insert_after(p2, "每天24:00的SOC与下一天0:00的SOC逐点相接，最大连接残差为0 kWh。结合母线平衡与状态递推残差检验，可以确认全年递归求解没有通过每日重置储能来虚假降低费用；各预测和情景参数也只使用决策日前已经完成的数据。")

    # 问题三：分别解释消融结果与代表日机制。
    p = find_exact(doc, "全年计算得到：仅在0:00制定计划的总费用为16391391.76元，采用0:00、6:00、12:00和18:00四次更新后的总费用为15731687.04元，调整净费用为499519.56元，紧急购电量为748796.98 kWh。图8-1、图8-2和表8-2—表8-4给出更新频率、代表日预测修正及指定日期的购电和储能结果。")
    set_text(p, "全年计算得到：仅在0:00制定计划的总费用为16391391.76元，采用0:00、6:00、12:00和18:00四次更新后的总费用为15731687.04元，调整净费用为499519.56元，紧急购电量为748796.98 kWh。图8-1比较更新频率的全年效果，图8-2展示四个代表日中各次预报如何修正尚未执行时段；表8-2—表8-4保留指定时点的精确计划和结算结果。")
    p = find_exact(doc, "与仅0:00计划相比，四次更新使总费用减少659704.73元，降幅约4.02%，紧急购电量下降28.03%。这说明新预报减少的高价应急损失超过了逐次调整成本，故四次更新在本组数据上具有实际收益。逐版本复算表明，每次增购、减购均相对上一版有效计划结算，已执行时段保持冻结，母线平衡、储电量边界和状态递推均满足约束。")
    set_text(p, "图8-1显示，更新节点由仅0时增加到0/6/12/18时后，全年总费用由1639.14万元降至1573.17万元，紧急购电量由1040.38 MWh降至748.80 MWh，分别下降4.02%和28.03%。前三次增加更新节点带来的费用降幅依次约为46.51万元、17.28万元和2.18万元，表明更新仍有收益，但边际改善逐步减小。")
    insert_after(p, "图8-2中，每条预测曲线只从其发布时刻开始，6时、12时和18时的预测利用新信息修正剩余轨迹，发布时刻之前的已执行区间保持不变。新预报减少的高价应急损失超过499519.56元的调整净费用，因此完整四次更新在本样本中值得保留；逐版本复算同时确认增购、减购均相对上一版有效计划结算，且母线平衡、SOC边界和状态递推全部满足约束。")

    # 问题四：把总成本层级和仅由价格未知产生的损失分开解释。
    p = find_exact(doc, "因果价格预测下，问题4-2与问题4-3的全年总费用分别为16124373.46元和16550967.30元；对应的完全价格信息反事实费用分别为15923572.41元和16401924.52元，全外生信息严格下界均为12780453.36元。图9-1和表9-2汇总不同信息条件下的费用及指定日期结果。")
    set_text(p, "因果价格预测下，问题4-2与问题4-3的全年总费用分别为16124373.46元和16550967.30元；对应的完全价格信息反事实费用分别为15923572.41元和16401924.52元，全外生信息严格下界均为12780453.36元。图9-1(a)比较三种信息条件下的成本层级，图9-1(b)单独放大价格未知造成的差额；表9-2给出指定日期的可核验结果。")
    p = find_exact(doc, "问题4-2和问题4-3相对完全价格信息基准分别增加200801.05元和149042.78元，反映价格未知带来的样本内代价。问题4-3的因果总费用比问题4-2高426593.84元，表明本样本中更频繁的调整虽然获得更新信息，但增减购电成本和剩余紧急购电仍超过其收益。严格下界仅用于刻画同时放宽未来价格、负荷、光伏信息和整数互斥后的理论边界，不作为现实可执行方案。")
    set_text(p, "图9-1(a)表明两种因果在线策略的费用均高于只放宽未来价格信息的反事实基准，也高于全外生信息下界。由于价格信息差仅占完全价格信息费用的约1%，单看总成本柱高不易辨认；图9-1(b)放大后可见，问题4-2和问题4-3因价格未知分别增加200801.05元和149042.78元，即1.261%和0.909%。")
    insert_after(p, "问题4-3的因果总费用仍比问题4-2高426593.84元，说明在本样本和题定结算规则下，更频繁的更新并不必然降低总成本：新增信息带来的收益被调整费用和剩余紧急购电抵消。全外生信息下界同时放宽价格、负荷和光伏未来信息，并放松整数互斥，只用于界定理论边界，不能作为现实可执行策略。")

    for table in doc.tables:
        format_table(table)

    # 统一所有图题、表题中的编号后空格及题注格式。
    for p in doc.paragraphs:
        text = p.text.strip()
        if re.match(r"^(图|表)\d+-\d+\S", text):
            text = re.sub(r"^((?:图|表)\d+-\d+)", r"\1 ", text)
            set_text(p, text, caption=True)
        elif re.match(r"^(图|表)\d+-\d+\s", text):
            normalize_caption(p)

    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category", "identifier"):
        try:
            setattr(props, attr, "")
        except Exception:
            pass

    doc.save(OUTPUT)
    AUDIT.write_text(
        json.dumps(
            {
                "source": str(INPUT.relative_to(ROOT)),
                "output": str(OUTPUT.relative_to(ROOT)),
                "replaced_figures": replaced,
                "added_figure": "图7-3 正式期跨日SOC连续性及日内运行范围",
                "table_count_formatted": len(doc.tables),
                "matlab_source_only": True,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(OUTPUT)
    print(f"replaced={len(replaced)} added=1 tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
