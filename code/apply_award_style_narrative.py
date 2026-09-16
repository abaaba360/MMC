"""按国奖论文的论证次序强化摘要、问题重述与问题分析，保持原排版基线。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_逻辑重构版.docx"
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_国奖行文强化版.docx"


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    hits = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        raise ValueError(f"段落定位失败: {prefix!r}, hits={len(hits)}")
    return hits[0]


def first_run_props(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def write_rich(paragraph: Paragraph, segments: list[tuple[str, bool]]) -> None:
    """替换段落正文并允许对指定片段直接加粗，段落格式不变。"""
    props = first_run_props(paragraph)
    ppr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not ppr:
            paragraph._p.remove(child)
    for text, bold in segments:
        run = paragraph.add_run(text)
        if props is not None:
            run._r.insert(0, deepcopy(props))
        run.bold = bold


def new_paragraph(template: Paragraph, segments: list[tuple[str, bool]], style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    if template._p.pPr is not None:
        element.append(deepcopy(template._p.pPr))
    paragraph = Paragraph(element, template._parent)
    if style:
        paragraph.style = style
    props = first_run_props(template)
    for text, bold in segments:
        run = paragraph.add_run(text)
        if props is not None:
            run._r.insert(0, deepcopy(props))
        run.bold = bold
    return paragraph


def insert_after(anchor: Paragraph, paragraph: Paragraph) -> Paragraph:
    anchor._p.addnext(paragraph._p)
    return paragraph


def append_after(anchor: Paragraph, template: Paragraph, text: str, *, bold_prefix: str = "", style: str | None = None) -> Paragraph:
    segments: list[tuple[str, bool]] = []
    if bold_prefix:
        if not text.startswith(bold_prefix):
            raise ValueError(f"文本未以加粗前缀开头: {bold_prefix}")
        segments.append((bold_prefix, True))
        segments.append((text[len(bold_prefix):], False))
    else:
        segments.append((text, False))
    return insert_after(anchor, new_paragraph(template, segments, style))


def remove_between(start: Paragraph, end: Paragraph) -> None:
    """删除两个标题之间的全部段落、表格和图片。"""
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def find_image_paragraph(doc: Document, target_ref: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        for rel_id in paragraph._p.xpath(".//a:blip/@r:embed"):
            if doc.part.rels[rel_id].target_ref == target_ref:
                return paragraph
    raise ValueError(f"找不到图片: {target_ref}")


def force_heading_bold(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            for run in paragraph.runs:
                run.bold = True


def anonymize_properties(doc: Document) -> None:
    """清除可能暴露作者、账号或队伍信息的核心文档属性。"""
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = ""
    props.subject = ""
    props.comments = ""
    props.keywords = ""
    props.category = ""
    props.identifier = ""


def main() -> None:
    doc = Document(INPUT)
    normal = find_paragraph(doc, "本文研究由光伏")
    heading2 = find_paragraph(doc, "2.1 从微网运行过程")

    # 摘要：总述后按四问分别形成“题意—模型—结果”闭环。
    write_rich(find_paragraph(doc, "本文研究由光伏"), [
        ("本文针对光伏出力、居民负荷与购电价格随时间变化条件下的微网经济调度问题，研究储能如何在满足供电平衡、容量与功率边界及充放电效率的前提下跨时段转移电量。四问的差异依次来自信息是否已知、计划能否更新以及偏差如何结算。因此，本文从实际能量流向出发，先建立每个10分钟时段的电量平衡和储电量递推关系，再按信息到达过程逐层构造可执行的调度模型。", False)
    ])
    write_rich(find_paragraph(doc, "对问题一"), [
        ("针对问题一：", True),
        ("全天电价、负荷与光伏曲线均已知，决策只需在144个时段间安排购电和储能。由母线守恒、储电量递推及容量功率边界建立确定性混合整数线性规划，并以二元状态排除同时充放电。", False),
        ("最优日购电量为59482.70 kWh，购电费用为35126.95元，较无储能基线降低26.90%", True),
        ("；连续松弛得到相同目标值，且能量平衡与状态转移残差均低于10⁻¹² kWh。", False),
    ])
    write_rich(find_paragraph(doc, "对问题二"), [
        ("针对问题二：", True),
        ("每日0:00只能依据历史数据制定全天计划，实际负荷与光伏在计划冻结后才揭示，由此形成严格的“先计划、后实现”两阶段结构。第一阶段确定正常购电与储能动作，第二阶段不再调整储能，仅按5倍电价紧急购电补足缺口；利用决策日前的整日净负荷残差构造场景，并以48小时延伸窗传递日末储能价值。", False),
        ("2月1日至12月31日总费用为15258022.89元，紧急购电量为404258.73 kWh", True),
        ("。", False),
    ])
    write_rich(find_paragraph(doc, "对问题三"), [
        ("针对问题三：", True),
        ("0:00、6:00、12:00和18:00的新光伏预报只改变未执行时段的信息，已发生的购电、充放电和费用不能回写。因此在每个发布时刻固定当前储电量，重算剩余时域并执行至下一次更新，形成滚动混合整数线性规划；增购、减购与最终紧急购电均按题定规则逐次记账。", False),
        ("四次更新的全年总费用为15731687.04元，较仅0:00计划减少659704.73元，紧急购电量下降28.03%", True),
        ("。", False),
    ])
    write_rich(find_paragraph(doc, "对问题四"), [
        ("针对问题四：", True),
        ("附件中的实时价格是事后实现值，决策时不能预先读取。本文仅用历史价格形成因果预测，分别驱动问题二和问题三，得到全年总费用", False),
        ("16124373.46元和16550967.30元", True),
        ("。为区分可执行结果与事后评价，另计算完全价格信息反事实基准及同时知道全年价格、负荷和光伏的严格下界12780453.36元。全部模型均通过跨日储电量、充放电互斥、费用账本和未来信息泄漏检查。", False),
    ])

    # 一、问题重述：按“背景—基本问题”组织，不提前堆模型名。
    h1 = find_paragraph(doc, "一、问题重述")
    h2 = find_paragraph(doc, "二、问题分析与模型路线")
    remove_between(h1, h2)
    cursor = h1
    cursor = append_after(cursor, heading2, "1.1 问题背景", style="Heading 2")
    cursor = append_after(cursor, normal,
        "在分布式光伏微网中，光伏可以直接向居民负荷供电，也可以在有富余时给储能充电；当光伏不足时，储能与外网共同补足负荷。由于光伏出力、居民负荷和电价均随时间变化，某一时段看似最便宜的动作未必对全天最有利。例如，当前保留的储电量可能在后续高价或高负荷时段产生更大价值，因此调度必须同时考虑当前供需与未来状态。")
    cursor = append_after(cursor, normal,
        "题目以10分钟为间隔给出价格、负荷、光伏及预测数据，并规定储能额定容量为12000 kWh、允许储电范围为1200—10800 kWh、最大充放电功率均为5000 kW，充放电效率为90%。储能使前后时段通过储电量相互关联，而不同问题又规定了不同的信息获得时刻和结算方式。故本题的关键不是人为设定光伏、储能、购电的固定优先级，而是在每个决策时刻只使用当时可得信息，安排满足物理约束且总费用尽可能低的能量流。")
    cursor = append_after(cursor, heading2, "1.2 基本问题", style="Heading 2")
    cursor = append_after(cursor, normal,
        "基于上述运行背景，本文需要利用附件数据依次解决以下四个相互递进的问题：")
    cursor = append_after(cursor, normal,
        "问题一：在某日144个时段的电价、居民负荷和光伏出力均已知的条件下，确定外网购电及储能充放电计划，使当日购电费用最小，并给出结果解释与可行性检验。",
        bold_prefix="问题一：")
    cursor = append_after(cursor, normal,
        "问题二：在每天0:00只能利用历史数据制定全天计划、实际负荷和光伏随后才揭示的条件下，确定正常购电和储能计划；若实际供电不足，仅以5倍实时电价紧急购电，并计算2月1日至12月31日总费用。",
        bold_prefix="问题二：")
    cursor = append_after(cursor, normal,
        "问题三：在问题二基础上利用0:00、6:00、12:00和18:00发布的光伏预报，对尚未执行时段滚动调整购电与储能计划；按增购、减购及紧急购电规则结算，并判断增加预测更新是否真正降低费用。",
        bold_prefix="问题三：")
    append_after(cursor, normal,
        "问题四：当外网电价改为事前未知的实时价格时，分别按照问题二和问题三的信息更新方式制定可执行策略，计算全年费用，并通过不同信息基准评价价格预测与滚动调整的作用。",
        bold_prefix="问题四：")

    # 二、问题分析：总述之后逐问解释“新增题意事实如何长成模型”。
    h2 = find_paragraph(doc, "二、问题分析与模型路线")
    h3 = find_paragraph(doc, "三、模型假设与数据处理")
    picture = find_image_paragraph(doc, "media/image1.png")
    caption = find_paragraph(doc, "图2-1")
    parent = picture._p.getparent()
    parent.remove(picture._p)
    caption._p.getparent().remove(caption._p)
    remove_between(h2, h3)
    write_rich(h2, [("二、问题分析", True)])
    cursor = h2
    cursor = append_after(cursor, normal,
        "本文要解决的是同一微网在四种信息条件下的调度问题。四问共用的物理层始终是“电量从哪里来、到哪里去，以及储电量怎样延续到下一时段”；真正随题目变化的是决策层的信息边界和结算层的费用规则。因而建模时先固定物理关系，再逐问识别在决策发生之前能够知道什么、之后允许修改什么，避免把四问写成彼此割裂的算法拼接。")
    cursor = append_after(cursor, heading2, "2.1 总体建模思路", style="Heading 2")
    cursor = append_after(cursor, normal,
        "对任一10分钟时段，外网购电、当期光伏和储能放电共同进入母线，居民负荷和储能充电从母线取电，无法利用的富余光伏记为弃光；充放电效率又决定下一时段储电量。这些题意事实直接给出每时段能量守恒、储电量递推、容量功率边界和充放电互斥关系。购电费用是购电量与价格的乘积，因此基本调度关系为线性时序优化，只有互斥状态需要二元变量表达。")
    cursor = append_after(cursor, normal,
        "在这一基础上，问题一给出全部日内信息；问题二使负荷与光伏在计划时未知，并限定事后只能紧急购电；问题三允许在新预报发布后修改未执行计划；问题四进一步使未来价格未知。模型由此按“完全已知—日前不确定—日内更新—价格不确定”的顺序递进，图2-1给出题目条件、数学结构与求解结果之间的对应关系。")
    cursor._p.addnext(picture._p)
    picture._p.addnext(caption._p)
    cursor = caption
    cursor = append_after(cursor, heading2, "2.2 问题一的分析", style="Heading 2")
    cursor = append_after(cursor, normal,
        "问题一中，144个时段的电价、负荷和光伏均在决策前已知，不存在预测误差和事后调整。需要确定的是每一时段购多少电、储能充多少或放多少，同时保证负荷得到满足、储电量不越界且日末回到初始水平。由于储电量把相邻时段串联，不能逐时贪心选择最便宜电源，而应将全天联立，以购电费用之和最小为目标。连续关系均为线性，加入充放电互斥二元变量后自然形成混合整数线性规划；连续松弛只用于检验本组数据下互斥约束是否影响最优值。")
    cursor = append_after(cursor, heading2, "2.3 问题二的分析", style="Heading 2")
    cursor = append_after(cursor, normal,
        "问题二与问题一的区别首先来自信息时序，而不是目标函数。每日0:00制定计划时，当天真实负荷和光伏尚未发生，故不能把附件中的当日实际值直接放入优化；计划冻结后，题目又只允许以5倍价格紧急购电补足缺口，不允许重新安排储能。因此决策必然分为两个阶段：第一阶段依据过去数据确定一份对所有可能实际情况都相同的购电与储能计划，第二阶段在实际净负荷揭示后仅计算紧急购电和弃光。为保留一天内误差的相关性，采用历史整日净负荷残差构造场景，并用期望总费用权衡正常购电与高价应急风险。")
    cursor = append_after(cursor, heading2, "2.4 问题三的分析", style="Heading 2")
    cursor = append_after(cursor, normal,
        "问题三明确在一天内四次提供新的光伏预报，因此全天计划不再需要冻结到底。每次更新时，过去动作已经执行、当前储电量已经形成，二者都不能改写；只有下一次更新之前及其后的未执行计划可以重算。这一“固定已执行状态—优化剩余时域—执行一段—接收新信息”的过程与题意完全对应，故采用确定性滚动优化。是否值得增加更新次数不能由模型名称判断，而应比较全年总费用、调整费用、紧急购电费用及紧急购电量；缺乏可识别参数和稳定样本外收益的场景风险增强不进入主模型。")
    cursor = append_after(cursor, heading2, "2.5 问题四的分析", style="Heading 2")
    append_after(cursor, normal,
        "问题四把原先已知的电价也改为实时变化量。附件给出的是事后完整价格路径，而现实决策只能使用当前及历史价格，因此首先要建立不读取未来值的因果价格预测，再分别嵌入问题二的日前两阶段模型和问题三的滚动模型。仅报告事后知道真实价格的最优结果会违反信息边界，所以本文把因果在线策略作为正式答案，同时设置完全价格信息反事实基准和全外生信息严格下界，用于区分价格未知造成的损失与整体预测、调度误差。")

    force_heading_bold(doc)
    anonymize_properties(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
