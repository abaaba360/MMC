"""在分问闭环版中补回国奖论文式第二章“问题分析”及总体思路图。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_分问闭环重构版.docx"
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_问题分析完善版.docx"
FLOWCHART = ROOT / "results" / "figures" / "paper_overall_flowchart.png"


def find(doc: Document, exact: str) -> Paragraph:
    hits = [p for p in doc.paragraphs if p.text.strip() == exact]
    if len(hits) != 1:
        raise ValueError(f"定位失败: {exact!r}, hits={len(hits)}")
    return hits[0]


def first_rpr(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def replace_text(paragraph: Paragraph, text: str, *, bold: bool = False) -> None:
    rpr = first_rpr(paragraph)
    ppr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not ppr:
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    run.bold = bold


def new_paragraph(template: Paragraph, text: str = "", style: str | None = None, *, bold: bool = False) -> Paragraph:
    elem = OxmlElement("w:p")
    if template._p.pPr is not None:
        elem.append(deepcopy(template._p.pPr))
    paragraph = Paragraph(elem, template._parent)
    if style:
        paragraph.style = style
    if text:
        run = paragraph.add_run(text)
        rpr = first_rpr(template)
        if rpr is not None:
            run._r.insert(0, rpr)
        run.bold = bold
    return paragraph


def insert_before(anchor: Paragraph, paragraph: Paragraph) -> Paragraph:
    anchor._p.addprevious(paragraph._p)
    return paragraph


def append_after(anchor: Paragraph, paragraph: Paragraph) -> Paragraph:
    anchor._p.addnext(paragraph._p)
    return paragraph


def renumber_figures_and_tables(doc: Document) -> None:
    chapter_map = {"5": "6", "6": "7", "7": "8", "8": "9"}
    pattern = re.compile(r"([图表])([5-8])-")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            original = run.text
            if not original:
                continue
            changed = pattern.sub(lambda m: f"{m.group(1)}{chapter_map[m.group(2)]}-", original)
            if changed != original:
                run.text = changed


def anonymize(doc: Document) -> None:
    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category", "identifier"):
        setattr(props, attr, "")


def main() -> None:
    if not FLOWCHART.exists():
        raise FileNotFoundError(FLOWCHART)
    doc = Document(INPUT)
    normal = next(p for p in doc.paragraphs if p.text.startswith("本文针对光伏出力"))
    h1_template = find(doc, "一、问题重述")
    h2_template = find(doc, "1.1 问题背景")
    assumptions = find(doc, "二、模型假设")

    # 先将原第二章及后续章节整体顺延，再插入新的问题分析章。
    heading_map = {
        "二、模型假设": "三、模型假设",
        "2.1 基本假设": "3.1 基本假设",
        "三、符号说明": "四、符号说明",
        "四、数据预处理与探索性分析": "五、数据预处理与探索性分析",
        "4.1 时间与单位对齐": "5.1 时间与单位对齐",
        "4.2 数据质量与基本特征": "5.2 数据质量与基本特征",
        "五、问题一的模型建立与求解": "六、问题一的模型建立与求解",
        "5.1 问题分析": "6.1 问题分析",
        "5.2 模型建立": "6.2 模型建立",
        "5.2.1 时段电量平衡": "6.2.1 时段电量平衡",
        "5.2.2 储能状态、边界与充放电互斥": "6.2.2 储能状态、边界与充放电互斥",
        "5.2.3 初始储电量条件": "6.2.3 初始储电量条件",
        "5.2.4 目标函数与完整模型": "6.2.4 目标函数与完整模型",
        "5.3 模型求解与计算结果": "6.3 模型求解与计算结果",
        "5.4 结果分析与检验": "6.4 结果分析与检验",
        "六、问题二的模型建立与求解": "七、问题二的模型建立与求解",
        "6.1 问题分析": "7.1 问题分析",
        "6.2 数据处理与情景构造": "7.2 数据处理与情景构造",
        "6.3 两阶段模型建立": "7.3 两阶段模型建立",
        "6.4 模型求解与计算结果": "7.4 模型求解与计算结果",
        "6.5 结果分析与检验": "7.5 结果分析与检验",
        "七、问题三的模型建立与求解": "八、问题三的模型建立与求解",
        "7.1 问题分析": "8.1 问题分析",
        "7.2 预测数据处理与状态更新": "8.2 预测数据处理与状态更新",
        "7.3 滚动优化模型建立": "8.3 滚动优化模型建立",
        "7.4 模型求解与计算结果": "8.4 模型求解与计算结果",
        "7.5 结果分析与检验": "8.5 结果分析与检验",
        "八、问题四的模型建立与求解": "九、问题四的模型建立与求解",
        "8.1 问题分析": "9.1 问题分析",
        "8.2 价格数据处理与因果预测": "9.2 价格数据处理与因果预测",
        "8.3 调度模型与评价基准": "9.3 调度模型与评价基准",
        "8.4 模型求解与计算结果": "9.4 模型求解与计算结果",
        "8.5 结果分析与信息价值": "9.5 结果分析与信息价值",
        "九、模型检验与结果可靠性": "十、模型检验与结果可靠性",
        "9.1 硬约束与费用账本检验": "10.1 硬约束与费用账本检验",
        "9.2 时间信息泄漏检验": "10.2 时间信息泄漏检验",
        "9.3 模型对照检验": "10.3 模型对照检验",
        "十、模型评价、改进与推广": "十一、模型评价、改进与推广",
        "10.1 模型优点": "11.1 模型优点",
        "10.2 模型局限与改进": "11.2 模型局限与改进",
        "10.3 模型推广": "11.3 模型推广",
        "十一、AI工具使用声明": "十二、AI工具使用声明",
        "十二、参考文献": "十三、参考文献",
        "十三、附录 完整源程序代码": "十四、附录 完整源程序代码",
    }
    # 先保存原始标题对象再统一改写，避免“5.1→6.1”后与原6.1重名。
    heading_targets = [(find(doc, old), new) for old, new in heading_map.items()]
    for paragraph, new in heading_targets:
        replace_text(paragraph, new, bold=True)
    renumber_figures_and_tables(doc)

    # 国奖论文式第二章：总述、逐问分析、总体思路图。
    h2 = insert_before(assumptions, new_paragraph(h1_template, "二、问题分析", "Heading 1", bold=True))
    cursor = append_after(h2, new_paragraph(normal,
        "本文研究由光伏、储能、外网和居民负荷构成的微网调度问题。四个问题使用相同的能量流和储能运行边界，但可获得的信息、计划允许调整的时刻以及偏差结算规则逐步变化。为保证各问衔接清楚，本文先识别每一问新增的题目条件，再据此确定数据处理、状态更新和优化求解步骤。"))

    cursor = append_after(cursor, new_paragraph(h2_template, "2.1 问题一的分析", "Heading 2", bold=True))
    cursor = append_after(cursor, new_paragraph(normal,
        "问题一要求在全天电价、负荷和光伏曲线均已知的情况下确定最低购电费用方案。首先把功率统一转换为10分钟时段电量，并定义购电量、充电量、放电量、弃光量和储电量；随后依据微网实际能量流建立每时段母线平衡、储电量递推、容量功率边界及充放电互斥约束；最后以全天购电费用最小为目标联立144个时段求解，并通过无储能基线、连续松弛和约束残差检验所得方案。"))

    cursor = append_after(cursor, new_paragraph(h2_template, "2.2 问题二的分析", "Heading 2", bold=True))
    cursor = append_after(cursor, new_paragraph(normal,
        "问题二在问题一基础上把当天实际负荷和光伏由事前已知改为计划制定时未知，并规定供电不足只能按5倍电价紧急购电。首先使用决策日前的历史数据预测当天净负荷，并按整日残差构造能够保留日内相关性的情景；随后把0:00确定的购电和储能动作作为第一阶段决策，把实际曲线揭示后的紧急购电和弃光作为第二阶段结果；最后按自然日期连续传递储电量，计算2月1日至12月31日的计划费用、紧急费用及总费用。"))

    cursor = append_after(cursor, new_paragraph(h2_template, "2.3 问题三的分析", "Heading 2", bold=True))
    cursor = append_after(cursor, new_paragraph(normal,
        "问题三进一步在0:00、6:00、12:00和18:00提供新的光伏预报，并允许调整尚未执行的计划。每次预报到达后，先固定已经执行的购电、充放电动作和当前储电量，再更新未来光伏与负荷输入，对剩余时域重新优化；随后只执行至下一次更新，并按照相邻两版有效计划计算增购与减购费用。最后比较不同更新频率下的总费用、调整费用和紧急购电量，判断日内更新是否产生实际收益。"))

    cursor = append_after(cursor, new_paragraph(h2_template, "2.4 问题四的分析", "Heading 2", bold=True))
    cursor = append_after(cursor, new_paragraph(normal,
        "问题四把外网电价也改为事前未知的实时价格。首先只使用决策时刻以前的历史价格建立因果预测，避免读取当天未来实现值；然后分别把预测价格嵌入问题二的0时两阶段调度和问题三的日内滚动调度，得到现实可执行费用；最后设置完全价格信息反事实基准和全外生信息严格下界，比较不同信息条件下的费用差异，从而评价价格预测和日内调整的价值。"))
    cursor = append_after(cursor, new_paragraph(normal,
        "综合以上分析，本文的总体求解过程如图2-1所示。"))
    image_p = new_paragraph(normal)
    image_p.alignment = 1
    image_p.add_run().add_picture(str(FLOWCHART), width=Inches(6.45))
    cursor = append_after(cursor, image_p)
    caption = new_paragraph(normal, "图2-1 文章总体思路图")
    caption.alignment = 1
    append_after(cursor, caption)

    anonymize(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
