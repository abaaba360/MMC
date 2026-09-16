from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"E:\微信聊天记录\xwechat_files\wxid_so1zh5t7c8rl22_8e76\msg\file\2026-09\C题论文_光伏微网储能购电滚动优化_全文格式与参考文献终修版(1)(1).docx")
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(r"D:\数模工作流\state\agent_outputs\new_c_paper_docx_audit.json")
ARTIFACT = Path(r"D:\数模工作流\state\agent_outputs\c_paper_template_artifact.md")


def iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def has_drawing(p: Paragraph) -> bool:
    return bool(p._p.xpath(".//w:drawing | .//w:pict"))


doc = Document(SOURCE)
blocks = []
for i, block in enumerate(iter_blocks(doc)):
    if isinstance(block, Paragraph):
        txt = block.text.strip()
        blocks.append({
            "i": i,
            "type": "p",
            "style": block.style.name if block.style else "",
            "text": txt,
            "drawing": has_drawing(block),
            "omml": len(block._p.xpath(".//m:oMath | .//m:oMathPara")),
        })
    else:
        preview = [[c.text.strip() for c in row.cells] for row in block.rows[:3]]
        blocks.append({"i": i, "type": "table", "rows": len(block.rows), "cols": len(block.columns), "preview": preview})

sections = []
for s in doc.sections:
    sections.append({
        "page_width_cm": round(s.page_width.cm, 3),
        "page_height_cm": round(s.page_height.cm, 3),
        "top_cm": round(s.top_margin.cm, 3),
        "bottom_cm": round(s.bottom_margin.cm, 3),
        "left_cm": round(s.left_margin.cm, 3),
        "right_cm": round(s.right_margin.cm, 3),
        "header_cm": round(s.header_distance.cm, 3),
        "footer_cm": round(s.footer_distance.cm, 3),
    })

payload = {
    "source": str(SOURCE),
    "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shape_count": len(doc.inline_shapes),
    "section_count": len(doc.sections),
    "sections": sections,
    "blocks": blocks,
}
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

styles = {}
for p in doc.paragraphs:
    if not p.text.strip() and not has_drawing(p):
        continue
    name = p.style.name if p.style else ""
    styles.setdefault(name, 0)
    styles[name] += 1

ARTIFACT.write_text(
    "\n".join([
        "# C题论文保真编辑契约",
        "",
        f"- 参考文件：`{SOURCE}`",
        f"- SHA-256：`{payload['sha256']}`",
        f"- 节数：{len(doc.sections)}；段落：{len(doc.paragraphs)}；表格：{len(doc.tables)}；内嵌图形：{len(doc.inline_shapes)}。",
        f"- 页面系统：{json.dumps(sections, ensure_ascii=False)}",
        f"- 段落样式分布：{json.dumps(styles, ensure_ascii=False)}",
        "- 内容流：摘要专页；问题重述；问题分析；模型假设；符号说明；数据预处理；问题一至问题四的模型建立与求解；模型检验；模型评价、改进与推广；AI工具使用声明；参考文献；附录。",
        "- 可编辑槽位：问题一至问题四中的数据结果图、图表引导段、图表后局部解读段，以及各问结果分析与检验正文；摘要仅在不损失原信息的前提下恢复完整性。",
        "- 必须保留：原正文、公式OMML、标题层级、页边距、页眉页脚、表格数据、参考文献、附录与源代码，除用户明确要求的图表叙事调整外不重写。",
        "- 图表规则：图前说明观察目的；图题置于图下；图后立即解释证据。表前说明观察目的；表题置于表上；表后立即解释证据。不得出现连续图表无正文间隔。",
        "- 结果分析规则：各问末节使用正常正文综合回答题目、解释经济或物理机制并说明验证结论，不重复充当图注合集。",
        "- 插图规则：仅使用本地MATLAB按真实求解结果生成的图；采用内嵌型，保留清晰PNG并控制在正文版心内。",
        "- 保真门禁：参考文件不修改；输出另存；节与页面几何保持；OMML数量不得下降；每页渲染检查；图表邻接规则结构化校验。",
    ]),
    encoding="utf-8",
)
print(OUT)
print(ARTIFACT)
print(json.dumps({k: payload[k] for k in ["sha256", "paragraph_count", "table_count", "inline_shape_count", "section_count"]}, ensure_ascii=False))
