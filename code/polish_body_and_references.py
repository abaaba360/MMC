"""统一 C 题论文正文段落格式并按 GB/T 7714 规范重排参考文献。"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_模板规范修订版.docx"
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_全文格式与参考文献终修版.docx"
AUDIT = ROOT / "state" / "agent_outputs" / "body_paragraph_format_audit.json"

REFERENCES = [
    "[1] 张亚琳. 具有电池储能系统的微电网分布式经济调度方案与SOC均衡控制[D]. 天津: 南开大学, 2025.",
    "[2] 李嘉伟, 巨云涛, 张璐, 等. 基于分布鲁棒模型预测控制的微电网多时间尺度优化调度[J]. 电力工程技术, 2024, 43(4): 45-55.",
    "[3] 张理, 王宝, 贾健雄, 等. 微电网功率预测与调度端到端协同优化方法[J]. 上海交通大学学报, 2025, 59(6): 720-731.",
    "[4] LIU C, QIN Y, ZHANG H. Real-time scheduling strategy for microgrids considering operation interval division of DGs and batteries[J]. Global Energy Interconnection, 2020, 3(5): 442-452.",
    "[5] OROZCO C, BORGHETTI A, DE SCHUTTER B, et al. Intra-day scheduling of a local energy community coordinated with day-ahead multistage decisions[J]. Sustainable Energy, Grids and Networks, 2022, 29: 100573.",
    "[6] ELKAZAZ M, SUMNER M, THOMAS D. Energy management system for hybrid PV-wind-battery microgrid using convex programming, model predictive and rolling horizon predictive control with experimental validation[J]. International Journal of Electrical Power & Energy Systems, 2020, 115: 105483.",
    "[7] 侯慧, 王晴, 薛梦雅, 等. 计及源荷不确定性及需求响应的离网型微电网两阶段日前经济调度[J]. 电力系统保护与控制, 2022, 50(13): 73-85.",
    "[8] 刘一欣, 郭力, 王成山. 微电网两阶段鲁棒优化经济调度方法[J]. 中国电机工程学报, 2018, 38(14): 4013-4022, 4307.",
    "[9] 王灿, 张雪菲, 凌凯, 等. 基于区间概率不确定集的微电网两阶段自适应鲁棒优化调度[J]. 中国电机工程学报, 2024, 44(5): 1750-1764.",
]


def find_index(doc: Document, text: str) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == text:
            return index
    raise ValueError(f"找不到段落：{text}")


def set_run_font(run, size: float = 12) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = None
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")


def direct_state(paragraph) -> dict[str, object]:
    pf = paragraph.paragraph_format
    return {
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "first_line_pt": round(pf.first_line_indent.pt, 2) if pf.first_line_indent else None,
        "left_pt": round(pf.left_indent.pt, 2) if pf.left_indent else None,
        "right_pt": round(pf.right_indent.pt, 2) if pf.right_indent else None,
        "space_before_pt": round(pf.space_before.pt, 2) if pf.space_before else None,
        "space_after_pt": round(pf.space_after.pt, 2) if pf.space_after else None,
        "line_spacing": str(pf.line_spacing),
    }


def is_display_math(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//m:oMathPara"))


def is_figure_holder(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def is_caption(paragraph) -> bool:
    # 图题/表题在编号后有空格；“图6-1显示……”属于正文，不能误判为题注。
    return bool(re.match(r"^(图|表)\s*\d+(?:[-－]\d+)?\s+", paragraph.text.strip()))


def remove_numbering(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def normalize_heading(paragraph, level: int) -> None:
    pf = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(12 if level == 1 else 6 if level == 2 else 3)
    pf.space_after = Pt(6 if level == 1 else 0)
    pf.keep_with_next = True
    pf.keep_together = True
    pf.page_break_before = False


def normalize_body(paragraph) -> None:
    remove_numbering(paragraph)
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(24)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.keep_with_next = False
    pf.keep_together = False
    pf.page_break_before = False
    pf.widow_control = True
    for run in paragraph.runs:
        set_run_font(run, 12)


def normalize_caption(paragraph) -> None:
    remove_numbering(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(run, 10.5)


def add_body_citations(doc: Document) -> list[str]:
    additions = [
        (
            "在分布式光伏微网中",
            " 储能经济调度、多时间尺度协调以及功率预测与调度协同等研究，为本文的模型构造提供了方法参考[1-3]。",
        ),
        (
            "每次滚动求解均继承问题一的能量平衡",
            " 该处理与微网日内调度和滚动时域控制的基本思想一致[4-6]。",
        ),
        (
            "在同一个决策日的所有可能实际情景下",
            " 这种将事前计划与事后补救分开的结构也常用于含源荷不确定性的两阶段微网调度[7-9]。",
        ),
    ]
    applied: list[str] = []
    for prefix, citation_sentence in additions:
        for paragraph in doc.paragraphs:
            if paragraph.text.strip().startswith(prefix) and citation_sentence.strip() not in paragraph.text:
                run = paragraph.add_run(citation_sentence)
                set_run_font(run, 12)
                applied.append(citation_sentence.strip())
                break
    return applied


def rewrite_references(doc: Document) -> None:
    ref_index = find_index(doc, "十三、参考文献")
    appendix_index = find_index(doc, "十四、附录")
    ref_heading = doc.paragraphs[ref_index]
    appendix = doc.paragraphs[appendix_index]

    for paragraph in list(doc.paragraphs[ref_index + 1:appendix_index]):
        paragraph._element.getparent().remove(paragraph._element)

    normalize_heading(ref_heading, 1)
    for reference in REFERENCES:
        paragraph = appendix.insert_paragraph_before(reference)
        paragraph.style = "Normal"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = paragraph.paragraph_format
        pf.left_indent = Pt(24)
        pf.first_line_indent = Pt(-24)
        pf.right_indent = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.keep_together = True
        pf.widow_control = True
        for run in paragraph.runs:
            set_run_font(run, 12)


def main() -> None:
    doc = Document(INPUT)
    start = find_index(doc, "一、问题重述")
    end = find_index(doc, "十三、参考文献")
    changed: list[dict[str, object]] = []

    for index, paragraph in enumerate(doc.paragraphs[start:end], start):
        text = paragraph.text.strip()
        if not text:
            continue
        before = direct_state(paragraph)
        style_name = paragraph.style.name
        if style_name.startswith("Heading "):
            normalize_heading(paragraph, int(style_name.rsplit(" ", 1)[1]))
        elif style_name == "Code" or is_display_math(paragraph) or is_figure_holder(paragraph):
            continue
        elif is_caption(paragraph):
            normalize_caption(paragraph)
        else:
            normalize_body(paragraph)
        after = direct_state(paragraph)
        if before != after:
            changed.append({
                "paragraph_index": index,
                "text": text[:100],
                "before": before,
                "after": after,
            })

    citations = add_body_citations(doc)
    rewrite_references(doc)

    # 二次扫描：保证正文不存在截图中所示的居中、悬挂缩进或异常段距。
    remaining_anomalies: list[dict[str, object]] = []
    start_after = find_index(doc, "一、问题重述")
    end_after = find_index(doc, "十三、参考文献")
    for index, paragraph in enumerate(doc.paragraphs[start_after:end_after], start_after):
        text = paragraph.text.strip()
        if (
            not text
            or paragraph.style.name.startswith("Heading ")
            or paragraph.style.name == "Code"
            or is_display_math(paragraph)
            or is_figure_holder(paragraph)
            or is_caption(paragraph)
        ):
            continue
        pf = paragraph.paragraph_format
        first_line = round(pf.first_line_indent.pt, 2) if pf.first_line_indent else 0.0
        left = round(pf.left_indent.pt, 2) if pf.left_indent else 0.0
        right = round(pf.right_indent.pt, 2) if pf.right_indent else 0.0
        before = round(pf.space_before.pt, 2) if pf.space_before else 0.0
        after = round(pf.space_after.pt, 2) if pf.space_after else 0.0
        if not (
            paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
            and first_line == 24.0
            and left == right == before == after == 0.0
            and pf.line_spacing_rule == WD_LINE_SPACING.ONE_POINT_FIVE
        ):
            remaining_anomalies.append({
                "paragraph_index": index,
                "text": text[:100],
                "state": direct_state(paragraph),
            })

    if remaining_anomalies:
        raise RuntimeError(f"仍有 {len(remaining_anomalies)} 个正文段落格式异常")

    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category", "identifier"):
        try:
            setattr(props, attr, "")
        except Exception:
            pass

    doc.save(OUTPUT)
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps({
        "source": str(INPUT.relative_to(ROOT)),
        "output": str(OUTPUT.relative_to(ROOT)),
        "changed_paragraph_count": len(changed),
        "citation_additions": citations,
        "reference_count": len(REFERENCES),
        "remaining_body_format_anomalies": remaining_anomalies,
        "changes": changed,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)
    print(
        f"normalized_paragraphs={len(changed)} references={len(REFERENCES)} "
        f"citations={len(citations)} remaining_anomalies={len(remaining_anomalies)}"
    )


if __name__ == "__main__":
    main()
