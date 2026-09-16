"""按《CUMCM2026 标准 Word 模板格式规范》修订 C 题论文。

本轮只重组论文表述与 Word 版式，不改变模型、公式和数值结果：
1. 将各问重复的“问题分析”集中到第二章并扩写；
2. 保留各问模型建立、求解、结果、分析与检验的现有闭环；
3. 重写模型评价、改进与推广；
4. 规范参考文献、表格和附录文件说明；
5. 保留全部 Word 原生 OMML 公式。
"""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_问题分析完善版.docx"
OUTPUT = ROOT / "delivery" / "C题论文_光伏微网储能购电滚动优化_模板规范修订版.docx"


def find(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"找不到段落：{text}")


def paragraph_after(doc: Document, paragraph: Paragraph) -> Paragraph | None:
    paragraphs = doc.paragraphs
    index = next(i for i, item in enumerate(paragraphs) if item._p is paragraph._p)
    return paragraphs[index + 1] if index + 1 < len(paragraphs) else None


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, *, east_asia: str = "宋体", latin: str = "Times New Roman",
                 size: float = 12, bold: bool = False) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = None
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_plain_text(paragraph: Paragraph, text: str, *, size: float = 12,
                   bold: bool = False, east_asia: str = "宋体") -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, east_asia=east_asia, size=size, bold=bold)


def new_paragraph_like(template: Paragraph, text: str = "", *, style: str | None = None,
                       bold: bool = False, east_asia: str = "宋体", size: float = 12) -> Paragraph:
    element = OxmlElement("w:p")
    if template._p.pPr is not None:
        element.append(deepcopy(template._p.pPr))
    paragraph = Paragraph(element, template._parent)
    if style:
        paragraph.style = style
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, east_asia=east_asia, size=size, bold=bold)
    return paragraph


def insert_after(anchor: Paragraph, paragraph: Paragraph) -> Paragraph:
    anchor._p.addnext(paragraph._p)
    return paragraph


def insert_before(anchor: Paragraph, paragraph: Paragraph) -> Paragraph:
    anchor._p.addprevious(paragraph._p)
    return paragraph


def format_body_paragraph(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, size=12, bold=bool(run.bold))


def replace_analysis(doc: Document, heading: str, paragraphs: list[str]) -> None:
    h = find(doc, heading)
    first = paragraph_after(doc, h)
    if first is None:
        raise ValueError(f"{heading} 后没有正文")
    set_plain_text(first, paragraphs[0])
    format_body_paragraph(first)
    cursor = first
    for text in paragraphs[1:]:
        added = new_paragraph_like(first, text)
        format_body_paragraph(added)
        cursor = insert_after(cursor, added)


def delete_section_intro(doc: Document, heading: str) -> None:
    h = find(doc, heading)
    following = paragraph_after(doc, h)
    if following is not None and not following.style.name.startswith("Heading"):
        remove_paragraph(following)
    remove_paragraph(h)


def rename_heading(doc: Document, old: str, new: str, level: int) -> None:
    p = find(doc, old)
    set_plain_text(p, new, size={1: 16, 2: 14, 3: 12}[level], bold=True, east_asia="黑体")
    p.style = f"Heading {level}"


def add_solver_paragraph(doc: Document, heading: str, text: str) -> None:
    h = find(doc, heading)
    following = paragraph_after(doc, h)
    if following is None:
        raise ValueError(f"{heading} 后没有段落")
    p = new_paragraph_like(following, text)
    format_body_paragraph(p)
    insert_after(h, p)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")
    for row_index, edges in ((0, (("top", 12), ("bottom", 6))),
                             (len(table.rows) - 1, (("bottom", 12),))):
        for edge_name, size in edges:
            for cell in table.rows[row_index].cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_borders = tc_pr.first_child_found_in("w:tcBorders")
                if tc_borders is None:
                    tc_borders = OxmlElement("w:tcBorders")
                    tc_pr.append(tc_borders)
                edge = tc_borders.find(qn(f"w:{edge_name}"))
                if edge is None:
                    edge = OxmlElement(f"w:{edge_name}")
                    tc_borders.append(edge)
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), str(size))
                edge.set(qn("w:color"), "000000")


def column_widths(column_count: int) -> list[float]:
    layouts = {
        2: [7.2, 8.8],
        3: [2.8, 9.7, 3.5],
        4: [3.0, 4.4, 4.3, 4.3],
        5: [2.8, 3.3, 3.3, 3.3, 3.3],
        6: [2.5, 2.7, 2.7, 2.7, 2.7, 2.7],
        7: [2.2, 2.3, 2.3, 2.3, 2.3, 2.3, 2.3],
        8: [2.0] + [2.0] * 7,
        9: [2.0] + [1.75] * 8,
        10: [2.0] + [1.55] * 9,
        11: [2.0] + [1.4] * 10,
    }
    return layouts.get(column_count, [16.0 / column_count] * column_count)


def format_tables(doc: Document) -> None:
    for table in doc.tables:
        columns = len(table.columns)
        widths = column_widths(columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        layout = table._tbl.tblPr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            table._tbl.tblPr.append(layout)
        layout.set(qn("w:type"), "fixed")
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_pr.append(repeat)
        font_size = 9 if columns >= 8 else (9.5 if columns >= 5 else 10.5)
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.width = Cm(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf = paragraph.paragraph_format
                    pf.first_line_indent = Pt(0)
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(run, size=font_size, bold=(row_index == 0))
        set_table_borders(table)


def format_captions(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not re.match(r"^(图|表)\d+(?:-\d+)?", text):
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = paragraph.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(4)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.0
        pf.keep_with_next = text.startswith("表")
        for run in paragraph.runs:
            set_run_font(run, size=10.5, bold=False)


def rewrite_abstract(doc: Document) -> None:
    title = doc.paragraphs[0]
    set_plain_text(title, "基于信息约束与非对称结算的光伏微网储能滚动优化研究",
                   size=16, bold=True, east_asia="黑体")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)

    intro = doc.paragraphs[2]
    set_plain_text(
        intro,
        "针对光伏出力、居民负荷与购电价格随时间变化条件下的微网经济调度问题，本文以每个10分钟时段的母线能量平衡和储电量状态转移为共同基础，重点研究信息到达时刻、计划调整权限与偏差结算规则对储能和购电决策的影响。",
        size=11,
    )
    intro.paragraph_format.first_line_indent = Pt(0)
    intro.paragraph_format.line_spacing = 1.25
    intro.paragraph_format.space_after = Pt(4)

    method = doc.paragraphs[3]
    clear_paragraph(method)
    parts = [
        ("针对问题一，", True),
        ("根据全天已知曲线建立含充放电互斥的确定性混合整数线性规划，最优日购电量为59482.70 kWh，费用为35126.95元，较无储能方案降低26.90%。", False),
        ("针对问题二，", True),
        ("依据“0时计划、事后仅紧急购电”的题意构造严格两阶段模型，利用历史整日净负荷残差形成场景；2月1日至12月31日总费用为15258022.89元，紧急购电量为404258.73 kWh。", False),
        ("针对问题三，", True),
        ("在四个光伏预报发布时刻冻结已执行决策并重算剩余时域，按增购、减购和紧急购电逐次结算；四次更新使全年费用减少659704.73元，紧急购电量下降28.03%。", False),
        ("针对问题四，", True),
        ("仅用历史价格构造因果预测并分别驱动日前与日内策略，全年费用为16124373.46元和16550967.30元；同时给出完全价格信息反事实基准和全外生信息严格下界12780453.36元。", False),
    ]
    for text, bold in parts:
        run = method.add_run(text)
        set_run_font(run, size=11, bold=bold)
    method.paragraph_format.first_line_indent = Pt(0)
    method.paragraph_format.line_spacing = 1.25
    method.paragraph_format.space_after = Pt(4)

    old = doc.paragraphs[4:7]
    conclusion = new_paragraph_like(
        method,
        "本文模型由能量流、信息流和费用流逐层推出，在不预设供电优先级的同时严格限制未来信息使用。约束残差、跨日状态、费用账本与对照方案的复核表明，所得策略具有物理可行性、可解释性和可复现性，可为含储能分布式微网的分时调度提供参考。",
        size=11,
    )
    conclusion.paragraph_format.first_line_indent = Pt(0)
    conclusion.paragraph_format.line_spacing = 1.25
    conclusion.paragraph_format.space_after = Pt(4)
    insert_after(method, conclusion)
    for paragraph in old:
        remove_paragraph(paragraph)


def rewrite_assumptions(doc: Document) -> None:
    texts = [
        "假设1（时段内平稳性）：每个10分钟时段内的负荷、光伏和电价均以附件给出的数值代表该时段平均水平。理由：附件的时间分辨率为10分钟，在该尺度内采用分段常值既与数据口径一致，也能直接完成电量换算。",
        "假设2（能量流完整性）：光伏可直接供给负荷，也可用于储能充电；模型不人为规定光伏、储能和外网的固定供电次序。理由：实际调度取决于当期供需、储能状态和未来电价，固定优先级会割裂跨时段决策。",
        "假设3（不向外网售电）：题目未给出反向上网通道和售电收益，富余光伏只能用于充电或弃置，弃光收益记为0。理由：这样处理不引入题目之外的交易机制。",
        "假设4（效率口径一致）：充电量和放电量均定义在交流母线侧，充、放电效率分别取题目给定的90%。理由：统一计量位置后，母线平衡与储电量递推之间不会重复计算损耗。",
        "假设5（跨日状态连续）：问题二至问题四的日末储电量直接作为次日初值；问题一仅采用题目明确给出的日初日末相等条件。理由：储能是连续运行设备，除题目另有规定外不应在每日零时人为复位。",
        "假设6（因果信息边界）：预测模型、标准化参数、情景残差及模型选择只使用决策时刻以前的数据，未来实际值仅用于计划冻结后的结算与评价。理由：该假设保证日前和滚动策略在现实中可执行。",
    ]
    heading = find(doc, "3.1 基本假设")
    cursor = paragraph_after(doc, heading)
    for text in texts:
        if cursor is None:
            raise ValueError("模型假设段落数量不足")
        next_p = paragraph_after(doc, cursor)
        cursor.style = doc.styles["Normal"]
        set_plain_text(cursor, text)
        format_body_paragraph(cursor)
        cursor.paragraph_format.first_line_indent = Pt(0)
        cursor = next_p


def rewrite_evaluation(doc: Document) -> None:
    start = find(doc, "11.1 模型优点")
    ai_heading = find(doc, "十二、AI工具使用声明")
    paragraphs = doc.paragraphs
    start_index = next(i for i, p in enumerate(paragraphs) if p._p is start._p)
    end_index = next(i for i, p in enumerate(paragraphs) if p._p is ai_heading._p)
    h2_template = start
    body_template = paragraphs[start_index + 1]
    for paragraph in paragraphs[start_index:end_index]:
        remove_paragraph(paragraph)

    blocks = [
        ("11.1 模型优点", [
            "（1）物理关系完整且解释明确。模型从母线能量守恒出发，以储电量递推连接相邻时段，并同时约束容量、充放电功率和工作状态；各变量均对应真实能量流，因而能够解释储能削峰、光伏消纳和弃光产生的原因。",
            "（2）信息边界与题意一致。问题二在0时冻结全天计划，实际发生后仅允许紧急购电；问题三只在新预报到达后修改尚未执行时段；问题四的价格预测也只使用历史数据，避免以事后信息改善事前策略。",
            "（3）费用结算具有可追溯性。正常购电、紧急购电、增购和减购分别记账，问题三和问题四的每次调整均相对上一版有效计划计算，因此总费用能够逐时段复算，且不会重复计费或遗漏退回金额。",
            "（4）模型简明且可稳定求解。除充放电互斥外，目标函数、状态方程和边界均为线性关系，形成混合整数线性规划；问题一连续松弛与整数模型同值，四问的平衡和状态残差均接近机器精度。",
        ]),
        ("11.2 模型缺点", [
            "（1）历史数据只覆盖一个完整年度，虽能反映日内规律和季节变化，但对跨年气候漂移、极端连续阴雨及异常价格尖峰的覆盖不足，历史残差情景可能低估样本外尾部风险。",
            "（2）题目未提供电池循环衰减、启停损耗、需量电费和上网交易参数，本文只能围绕购电与偏差费用优化；当这些工程成本不可忽略时，当前目标函数对长期运行经济性的描述仍不完整。",
        ]),
        ("11.3 模型改进", [
            "针对样本覆盖不足，可在获得多年气象和运行数据后，按季节、天气型和预测提前量分别估计条件残差分布，并采用滚动回测确定情景数量及尾部保留比例；当极端风险的样本外费用显著下降时，再引入机会约束或条件风险价值项。",
            "针对设备成本缺项，可依据电池厂家给出的循环寿命曲线建立分段线性衰减费用，并把启停损耗、最大需量费用及允许上网时的售电收益加入费用账本；新增参数应先进行可识别性与敏感性检验，避免用主观权重替代真实工程数据。",
        ]),
        ("11.4 模型推广", [
            "本文框架可推广至含风电、可控负荷、电动汽车和多类储能的园区微网。推广时，在母线平衡中增加相应能量流，在状态方程中补充设备运行特性，并按当地市场规则重写购售电和偏差结算项即可；“已执行决策不可回写、预测只使用当时可得数据”的因果边界仍应保持不变。",
        ]),
    ]
    for heading_text, body_texts in blocks:
        heading = new_paragraph_like(h2_template, heading_text, style="Heading 2",
                                     bold=True, east_asia="黑体", size=14)
        insert_before(ai_heading, heading)
        for text in body_texts:
            body = new_paragraph_like(body_template, text)
            format_body_paragraph(body)
            insert_before(ai_heading, body)


def rewrite_references(doc: Document) -> None:
    ref_heading = find(doc, "十三、参考文献")
    appendix = find(doc, "十四、附录 完整源程序代码")
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p._p is ref_heading._p) + 1
    end = next(i for i, p in enumerate(paragraphs) if p._p is appendix._p)
    old = paragraphs[start:end]
    full_text = " ".join(p.text.strip() for p in old if p.text.strip())
    refs = [x.strip() for x in re.split(r"(?=\[\d+\])", full_text) if x.strip()]
    template = old[0]
    for p in old:
        remove_paragraph(p)
    for ref in refs:
        p = new_paragraph_like(template, ref, size=10.5)
        # 英文题名与网址在两端对齐时容易被拉出过大的词间距，参考文献统一左对齐。
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.left_indent = Pt(21)
        pf.first_line_indent = Pt(-21)
        pf.line_spacing = 1.2
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        insert_before(appendix, p)


def expand_appendix(doc: Document) -> None:
    appendix = find(doc, "十四、附录 完整源程序代码")
    set_plain_text(appendix, "十四、附录", size=16, bold=True, east_asia="黑体")
    intro = paragraph_after(doc, appendix)
    set_plain_text(intro, "附录给出提交文件清单、程序运行说明及与支撑材料一致的完整源程序。")
    format_body_paragraph(intro)
    intro.paragraph_format.first_line_indent = Pt(0)

    first_code = find(doc, "A.1common_milp.py")
    h2_template = first_code
    body_template = paragraph_after(doc, first_code)
    h = new_paragraph_like(h2_template, "A.1 附件与支撑材料文件列表", style="Heading 2",
                           bold=True, east_asia="黑体", size=14)
    insert_before(first_code, h)
    descriptions = [
        "（1）result1.xlsx、result2.xlsx、result3.xlsx、result4-2.xlsx和result4-3.xlsx：按题目附件5格式填写的五份结果工作簿。",
        "（2）code文件夹：包含统一运行入口run_all.py、四问求解程序、因果预测、结果复核、图表生成和结果导出程序。全部程序使用Python编写。",
        "（3）results文件夹：包含四问关键结果、逐日汇总、决策明细及论文所用图表，可用于复核论文中的数值和结论。",
        "（4）AI工具使用详情.pdf：记录竞赛过程中AI工具的名称、版本、使用环节、主要交互、采纳修改及人工核验情况。",
        "程序运行环境为Python 3.11及以上版本，统一入口为run_all.py。代码编写与调试过程中使用了OpenAI Codex，模型结构、参数口径、结果与结论由本队结合题目和程序输出复核。",
    ]
    for text in descriptions:
        p = new_paragraph_like(body_template, text)
        format_body_paragraph(p)
        p.paragraph_format.first_line_indent = Pt(0)
        insert_before(first_code, p)

    code_headings = [p for p in doc.paragraphs if p.style.name == "Heading 2" and re.match(r"A\.\d+", p.text.strip()) and "文件列表" not in p.text]
    roles = {
        "common_milp.py": "公共混合整数线性规划结构与结果对象",
        "q1_model.py": "问题一确定性调度、连续松弛与无储能对照",
        "formal_data_loader.py": "附件数据读取、字段校验与单位转换",
        "causal_forecasts.py": "因果负荷、光伏和价格预测及残差情景构造",
        "q2_model.py": "问题二严格日前两阶段调度",
        "q34_helpers.py": "问题三和问题四的滚动状态与费用结算函数",
        "q3_model.py": "问题三预报更新频率消融与全年滚动回测",
        "q4_model.py": "问题四因果价格策略和信息基准计算",
        "verify_results.py": "能量平衡、状态递推、互斥和费用账本复核",
        "export_result_workbooks.py": "向附件5结果工作簿回填计算结果",
        "visualize_results.py": "根据已验证结果生成论文图表",
    }
    for index, heading in enumerate(code_headings, 2):
        original_name = re.sub(r"^A\.\d+\s*", "", heading.text.strip())
        name_match = re.search(r"([A-Za-z_][A-Za-z0-9_]*\.py)", original_name)
        if not name_match:
            continue
        name = name_match.group(1)
        set_plain_text(heading, f"A.{index} {name}", size=14, bold=True, east_asia="黑体")
        note = new_paragraph_like(
            body_template,
            f"语言：Python；作用：{roles.get(name, '论文计算与结果整理')}；具体输入输出见程序注释与支撑材料README。",
        )
        format_body_paragraph(note)
        note.paragraph_format.first_line_indent = Pt(0)
        insert_after(heading, note)


def normalize_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Code":
            continue
        if paragraph.style.name.startswith("Heading"):
            level = int(paragraph.style.name[-1])
            for run in paragraph.runs:
                set_run_font(run, east_asia="黑体", size={1: 16, 2: 14, 3: 12}[level], bold=True)


def main() -> None:
    doc = Document(INPUT)
    rewrite_abstract(doc)

    replace_analysis(doc, "2.1 问题一的分析", [
        "问题一给出了同一典型日内144个时段的电价、居民负荷和光伏出力，要求安排外网购电及储能充放电，使全天购电费用最低。由于功率数据以千瓦给出，而费用和储能状态按千瓦时计算，首先必须把每个10分钟时段的功率统一转换为电量，明确所有决策量的计量位置和时间含义。",
        "从能量流向看，光伏可直接供给负荷，也可在有富余时给储能充电；外网购电和储能放电用于补足母线缺口，无法消纳的光伏则形成弃光。由此需要建立逐时段母线平衡，并用储电量递推把当前充放电与下一时段状态连接起来，同时加入容量、功率和充放电互斥边界。全天曲线在决策前均已知，目标费用和全部约束均可一次列出，因此采用确定性混合整数线性规划求解，并用无储能方案、连续松弛和约束残差进行检验。",
    ])
    replace_analysis(doc, "2.2 问题二的分析", [
        "问题二的设备结构与问题一相同，但决策时点发生了根本变化：每天0:00制定全天计划时，当天真实负荷和光伏尚未出现，实际供电不足只能在计划冻结后按5倍价格紧急购电。因此，模型不能把附件中的当天真实曲线直接当作已知量，而要区分计划阶段能够使用的历史信息与实现阶段才观察到的实际结果。",
        "为描述全天计划面临的不确定性，先预测当天净负荷，再从决策日前的整日预测残差中抽取场景，使144个时段的共同偏移和峰谷变化得以保留。0时确定的正常购电和储能动作在所有场景中保持一致，实际曲线揭示后仅允许紧急购电和弃光随情景变化，由题意自然形成严格两阶段随机规划。全年计算还必须逐日传递储电量，并通过延伸求解时域保留日末电量的后续价值。",
    ])
    replace_analysis(doc, "2.3 问题三的分析", [
        "问题三在问题二基础上增加了0:00、6:00、12:00和18:00四次光伏预报。新预报到达时，过去时段的购电、充放电和费用已经实现，当前储电量也已经形成；这些量不能重新修改，只有尚未执行时段的计划能够利用新信息调整。因此，求解过程必须显式保存每个更新时点的已执行决策与当前状态。",
        "每次更新时，以当前储电量为新初值，替换剩余时域的光伏预测并修正负荷预测，再重新优化尚未执行部分，执行到下一次预报发布后重复上述过程，由此形成确定性滚动优化。费用不能只按最终计划计算，而要将新计划与上一版有效计划逐时段比较，分别记录增购、减购退回和最终紧急购电。最后比较不同更新频率下的费用与紧急购电量，判断新增预报是否具有实际经济价值。",
    ])
    replace_analysis(doc, "2.4 问题四的分析", [
        "问题四进一步把外网电价由事前已知的分时价格改为实时价格。附件4给出的是最终实现路径，但正式决策时只能看到当前及历史价格，若直接读取当天未来价格，就会把事后信息错误地带入事前方案。因此，需要为每个0时计划或日内更新时点建立相应的历史数据截断，并只用截断日前的数据选择预测方法和参数。",
        "在得到因果价格预测后，分别嵌入问题二的日前两阶段模型和问题三的日内滚动模型，计划仍按当时预测价格制定，费用则按实际价格事后结算。为区分价格未知造成的损失与负荷、光伏等其他不确定性，再设置只放宽价格信息的完全价格信息反事实基准，以及同时知道价格、负荷和光伏的全外生信息下界，从而形成现实策略与信息价值评价的统一比较框架。",
    ])

    rewrite_assumptions(doc)

    for heading in ("6.1 问题分析", "7.1 问题分析", "8.1 问题分析", "9.1 问题分析"):
        delete_section_intro(doc, heading)
    renames = [
        ("6.2 模型建立", "6.1 模型建立", 2),
        ("6.2.1 时段电量平衡", "6.1.1 时段电量平衡", 3),
        ("6.2.2 储能状态、边界与充放电互斥", "6.1.2 储能状态、边界与充放电互斥", 3),
        ("6.2.3 初始储电量条件", "6.1.3 初始储电量条件", 3),
        ("6.2.4 目标函数与完整模型", "6.1.4 目标函数与完整模型", 3),
        ("6.3 模型求解与计算结果", "6.2 模型求解与计算结果", 2),
        ("6.4 结果分析与检验", "6.3 结果分析与检验", 2),
        ("7.2 数据处理与情景构造", "7.1 数据处理与情景构造", 2),
        ("7.3 两阶段模型建立", "7.2 两阶段模型建立", 2),
        ("7.4 模型求解与计算结果", "7.3 模型求解与计算结果", 2),
        ("7.5 结果分析与检验", "7.4 结果分析与检验", 2),
        ("8.2 预测数据处理与状态更新", "8.1 预测数据处理与状态更新", 2),
        ("8.3 滚动优化模型建立", "8.2 滚动优化模型建立", 2),
        ("8.4 模型求解与计算结果", "8.3 模型求解与计算结果", 2),
        ("8.5 结果分析与检验", "8.4 结果分析与检验", 2),
        ("9.2 价格数据处理与因果预测", "9.1 价格数据处理与因果预测", 2),
        ("9.3 调度模型与评价基准", "9.2 调度模型与评价基准", 2),
        ("9.4 模型求解与计算结果", "9.3 模型求解与计算结果", 2),
        ("9.5 结果分析与信息价值", "9.4 结果分析与信息价值", 2),
    ]
    for old, new, level in renames:
        rename_heading(doc, old, new, level)

    add_solver_paragraph(doc, "6.2 模型求解与计算结果", "将144个时段的购电、充电、放电、弃光、储电量及充放电状态按时间顺序组成决策向量，把式（1）—（6）写成稀疏线性约束，采用HiGHS分支定界算法求解混合整数线性规划。求解完成后按原始时间顺序还原各能量流，并逐时复算平衡、状态和互斥约束；同时求解无储能基线与连续松弛模型，用于评价储能收益和整数约束的实际影响。")
    add_solver_paragraph(doc, "7.3 模型求解与计算结果", "全年按自然日期顺序递归求解。每个决策日先用截止前一日的数据生成净负荷基线和历史残差场景，再求解48小时两阶段模型，仅执行并结算前24小时，随后把日末储电量传给下一日。程序分别累加计划购电费、紧急购电费和紧急购电量，并对指定日期输出逐时段计划、实际缺口与储能状态。")
    add_solver_paragraph(doc, "8.3 模型求解与计算结果", "对每个运行日依次设置0时、6时、12时和18时四个决策节点。每到一个节点，程序读取当前储电量和最新预测，冻结此前已经执行的变量，只对剩余时段重新求解；新旧有效购电计划的差额随即进入调整账本。为检验更新频率的价值，分别计算仅0时、增加6时、增加12时以及四次更新的完整年度结果。")
    add_solver_paragraph(doc, "9.3 模型求解与计算结果", "每个价格预测原点先截断历史样本并选择预测误差较低的方法，再把预测价格传入相应调度模型。问题4-2每日只在0时求解一次，问题4-3随四次光伏预报同步更新；两者均以实际价格完成事后结算。另在保持其余条件不变时放宽未来价格信息形成反事实基准，并一次性输入全部外生实际路径求得严格下界。")

    rewrite_evaluation(doc)
    rewrite_references(doc)
    expand_appendix(doc)
    normalize_styles(doc)
    format_tables(doc)
    format_captions(doc)

    # 清除问题四评价基准条目遗留的正文加粗，保持正文层级一致。
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith((
            "（1）完全价格信息反事实基准",
            "（2）全外生信息下界",
        )):
            for run in paragraph.runs:
                run.bold = False

    # 问题一结果图略微缩小，使其与前置公式及求解说明落在同一页，减少无意义留白。
    if len(doc.inline_shapes) > 1:
        q1_figure = doc.inline_shapes[1]
        aspect_ratio = q1_figure.height / q1_figure.width
        q1_figure.width = Cm(12.5)
        q1_figure.height = int(q1_figure.width * aspect_ratio)

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    props = doc.core_properties
    for attr in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category", "identifier"):
        try:
            setattr(props, attr, "")
        except Exception:
            pass
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
