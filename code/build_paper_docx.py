"""生成可编辑、公式为原生OMML的CUMCM论文DOCX。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "论文_C题_第一版.md"
OUTPUT = ROOT / "paper" / "C题论文_第一版.docx"
XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
CODE_FILES = [
    "common_milp.py", "q1_model.py", "formal_data_loader.py", "causal_forecasts.py",
    "q2_model.py", "q34_helpers.py", "q3_model.py", "q4_model.py",
    "verify_results.py", "export_result_workbooks.py", "visualize_results.py",
]
TRANSFORM = etree.XSLT(etree.parse(str(XSL)))


def set_run_font(run, east="宋体", latin="Times New Roman", size=12, bold=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_math(paragraph, latex: str):
    latex = re.sub(r"\\tag\{.*?\}", "", latex).strip()
    mathml = latex_to_mathml(latex)
    omml = TRANSFORM(etree.fromstring(mathml.encode("utf-8")))
    paragraph._p.append(parse_xml(etree.tostring(omml.getroot())))


def add_inline(paragraph, text: str, size=12, bold=False):
    parts = re.split(r"(\$[^$]+\$|\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("$") and part.endswith("$"):
            add_math(paragraph, part[1:-1])
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, east="等线", latin="Consolas", size=max(9, size - 1))
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold)


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, attrs in edges.items():
        tag = "w:" + edge
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key, value in attrs.items():
            element.set(qn("w:" + key), str(value))


def configure_section(section, first=False):
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    section.different_first_page_header_footer = first


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    section._sectPr.append(pg)


def add_table(doc, rows):
    parsed = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    table = doc.add_table(rows=len(parsed), cols=len(parsed[0]))
    table.autofit = True
    table.alignment = 1
    for r, values in enumerate(parsed):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 or c > 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=10.5, bold=(r == 0))
            cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd")) if False else None
            set_cell_border(cell, top={"val": "nil"}, left={"val": "nil"},
                            right={"val": "nil"}, bottom={"val": "nil"})
    for c in range(len(parsed[0])):
        set_cell_border(table.cell(0, c), top={"val": "single", "sz": "12", "color": "000000"},
                        bottom={"val": "single", "sz": "6", "color": "000000"})
        set_cell_border(table.cell(len(parsed)-1, c), bottom={"val": "single", "sz": "12", "color": "000000"})
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def add_heading(doc, title, level):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    title = re.sub(r"^\d+(?:\.\d+)*\s*", "", title)
    add_inline(p, title, size={1:16, 2:14, 3:12}.get(level, 12), bold=True)
    return p


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_section(doc.sections[0], first=True)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Pt(24)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    for level, size in ((1,16),(2,14),(3,12)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True
    if "Code" not in [s.name for s in doc.styles]:
        code_style = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        code_style.font.size = Pt(8)
        code_style.paragraph_format.first_line_indent = Pt(0)
        code_style.paragraph_format.line_spacing = 1.0

    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    add_inline(p, title, size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, "摘    要", size=16, bold=True)
    abs_start = lines.index("## 摘要") + 1
    body_start = next(i for i, line in enumerate(lines) if line.startswith("## 1 "))
    # 摘要须与标题、关键词同处第一页（摘要专用页不超过1页），
    # 故摘要块不设首行缩进并收紧行距与字号。
    for line in lines[abs_start:body_start]:
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, line, size=11)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    add_page_number(body_section)

    i = body_start
    appendix_seen = False
    while i < len(lines):
        text = lines[i].strip()
        if not text:
            i += 1
            continue
        if text.startswith("## 附录"):
            appendix_seen = True
            break
        if text.startswith("## "):
            add_heading(doc, text[3:], 1); i += 1; continue
        if text.startswith("### "):
            add_heading(doc, text[4:], 2); i += 1; continue
        if text.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", text)
            if m:
                image = ROOT / "results" / "figures" / (Path(m.group(2)).stem + ".png")
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image), width=Cm(14.0))
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.first_line_indent = Pt(0)
                add_inline(cap, m.group(1), size=10.5)
            i += 1; continue
        if text.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i+1]):
            rows = [lines[i]]; i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(doc, rows); continue
        if text.startswith("$$"):
            eq = []
            if text.endswith("$$") and len(text) > 4:
                eq = [text[2:-2].strip()]; i += 1
            else:
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("$$"):
                    eq.append(lines[i]); i += 1
                i += 1
            full = " ".join(eq)
            tag = re.search(r"\\tag\{(.*?)\}", full)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            add_math(p, full)
            if tag:
                add_inline(p, f"    ({tag.group(1)})")
            continue
        if re.match(r"^\d+\.\s+", text):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                p = doc.add_paragraph(style="List Number")
                add_inline(p, re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            continue
        para = [text]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{2,3}\s|!\[|\||\$\$|\d+\.\s)", lines[i].strip()):
            para.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); add_inline(p, " ".join(para))

    if appendix_seen:
        doc.add_section(WD_SECTION.NEW_PAGE)
        add_heading(doc, "附录 完整源程序代码", 1)
        p = doc.add_paragraph(); add_inline(p, "以下程序与支撑材料中的同名源文件完全一致。")
        for name in CODE_FILES:
            add_heading(doc, name, 2)
            content = (ROOT / "code" / name).read_text(encoding="utf-8")
            for line in content.splitlines():
                p = doc.add_paragraph(style="Code")
                p.add_run(line or " ")

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = title
    doc.core_properties.subject = ""
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
