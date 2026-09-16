from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(r"E:\微信聊天记录\xwechat_files\wxid_so1zh5t7c8rl22_8e76\msg\file\2026-09\C题论文_光伏微网储能购电滚动优化_全文格式与参考文献终修版(1)(1).docx")
OUTPUT = Path(r"D:\数模工作流\delivery\C题论文_光伏微网储能购电滚动优化_图表逐项解读版.docx")
REPORT = Path(r"D:\数模工作流\state\agent_outputs\c_paper_figure_table_narrative_validation.json")


def blocks(doc):
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.append(Paragraph(child, doc))
        elif child.tag == qn("w:tbl"):
            out.append(Table(child, doc))
    return out


def drawing(p):
    return isinstance(p, Paragraph) and bool(p._p.xpath(".//w:drawing | .//w:pict"))


def text(b):
    return b.text.strip() if isinstance(b, Paragraph) else ""


s = Document(SOURCE)
d = Document(OUTPUT)
bb = blocks(d)
checks = []

checks.append(("source_unchanged", hashlib.sha256(SOURCE.read_bytes()).hexdigest() == "a4b82149f8266f529aa5785bde42b83e5481a966e559a23ad404d996d3395e27"))
checks.append(("abstract_verbatim", [p.text for p in s.paragraphs[:8]] == [p.text for p in d.paragraphs[:8]]))
checks.append(("section_count_preserved", len(s.sections) == len(d.sections) == 2))
checks.append(("table_count_preserved", len(s.tables) == len(d.tables) == 14))
checks.append(("figure_count_expected", len(d.inline_shapes) == 8))

src_omml = len(s.element.xpath(".//m:oMath | .//m:oMathPara"))
out_omml = len(d.element.xpath(".//m:oMath | .//m:oMathPara"))
checks.append(("omml_preserved", out_omml == src_omml))

figure_caps = [
    "图6-1 典型日功率、储能动作与SOC轨迹",
    "图7-1 四个代表日的日前预测、实际净负荷与紧急购电",
    "图7-2 正式期月度费用分解与紧急购电量",
    "图7-3 正式期跨日SOC连续性及日内运行范围",
    "图8-1 不同日内更新频率的费用与紧急购电量对比",
    "图8-2 四个指定日的滚动预测修正效果",
    "图9-1 因果策略、完全价格信息基准与全信息下界的费用对比",
]
table_caps = [
    "表6-1 问题一全天调度指标与对照", "表6-2 指定时段购电量及全天汇总", "表6-3 储能分段充/放电量及首末储电量",
    "表7-1 指定日期、指定时段的购电量及全天汇总", "表7-2 指定日期储能分段充/放电量（单元格为充电量/放电量，kWh）", "表7-3 指定日期紧急购电",
    "表8-1 不同预报更新频率的全年代价对照", "表8-2 指定日期、指定时段的0时计划量/最终有效量及全天汇总", "表8-3 指定日期储能分段充/放电量（单元格为充电量/放电量，kWh）", "表8-4 指定日期紧急购电",
    "表9-1 波动电价下三类策略的费用对照", "表9-2 波动电价下指定日期结果",
]

contexts = {}
for cap in figure_caps:
    ids = [i for i, b in enumerate(bb) if text(b) == cap]
    ok = len(ids) == 1
    if ok:
        i = ids[0]
        ok = i >= 2 and i + 1 < len(bb) and drawing(bb[i-1]) and isinstance(bb[i-2], Paragraph) and bool(text(bb[i-2])) and isinstance(bb[i+1], Paragraph) and bool(text(bb[i+1]))
        contexts[cap] = [text(bb[i-2]), "[DRAWING]", text(bb[i]), text(bb[i+1])]
    checks.append((f"figure_local_narrative::{cap}", ok))

for cap in table_caps:
    ids = [i for i, b in enumerate(bb) if text(b) == cap]
    ok = len(ids) == 1
    if ok:
        i = ids[0]
        ok = i >= 1 and i + 2 < len(bb) and isinstance(bb[i-1], Paragraph) and bool(text(bb[i-1])) and isinstance(bb[i+1], Table) and isinstance(bb[i+2], Paragraph) and bool(text(bb[i+2]))
        contexts[cap] = [text(bb[i-1]), text(bb[i]), "[TABLE]", text(bb[i+2])]
    checks.append((f"table_local_narrative::{cap}", ok))

for heading, nxt in [
    ("6.3 结果分析与检验", "七、问题二的模型建立与求解"),
    ("7.4 结果分析与检验", "八、问题三的模型建立与求解"),
    ("8.4 结果分析与检验", "九、问题四的模型建立与求解"),
    ("9.4 结果分析与信息价值", "十、模型检验与结果可靠性"),
]:
    a = next(i for i, b in enumerate(bb) if text(b) == heading)
    z = next(i for i, b in enumerate(bb) if text(b) == nxt)
    inner = [b for b in bb[a+1:z] if isinstance(b, Paragraph) and text(b)]
    ok = len(inner) >= 2 and all(not text(p).startswith(("图", "表")) for p in inner)
    checks.append((f"synthesis_prose::{heading}", ok))

passed = all(v for _, v in checks)
payload = {
    "passed": passed,
    "source_omml": src_omml,
    "output_omml": out_omml,
    "checks": [{"name": k, "passed": v} for k, v in checks],
    "contexts": contexts,
}
REPORT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"passed": passed, "checks": len(checks), "failed": [k for k, v in checks if not v], "source_omml": src_omml, "output_omml": out_omml}, ensure_ascii=False))
print(REPORT)
raise SystemExit(0 if passed else 1)
