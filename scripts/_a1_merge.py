# -*- coding: utf-8 -*-
"""A题：组装论文 = 前四部分(final_paper) + 第五部分(整合严谨版)。
删除 final_paper 旧第五部分(62-190)与旧六/七(191-226)，插入整合严谨版元素0-141。
两文件16张图片rId映射完全一致，无需重映射。
"""
import sys, io, shutil, copy, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

BASE = r'E:\微信聊天记录\xwechat_files\wxid_so1zh5t7c8rl22_8e76\msg\file\2026-08\final_paper_A题标题黑体统一版(1).docx'
SRC  = r'E:\微信聊天记录\xwechat_files\wxid_so1zh5t7c8rl22_8e76\msg\file\2026-08\选题A_问题一至问题三_整合严谨版(1).docx'
OUT  = r'd:\数模工作流\delivery\A题_智能评估\A题_数学建模论文智能评估_摘要参考文献版.docx'

os.makedirs(os.path.dirname(OUT), exist_ok=True)
shutil.copy2(BASE, OUT)
os.chmod(OUT, 0o666)
print('已复制基座 ->', OUT)

doc = Document(OUT)
body = doc.element.body
els = list(body)
assert len(els) == 263, 'final_paper body len=%d 期望263' % len(els)

anchor = els[61]  # 五、模型建立与求解 引言段，之后插入新第五部分
for i in range(62, 227):   # 删除 62..226（旧五/六/七）
    body.remove(els[i])
print('已删除旧第五部分及旧六/七 (元素62-226)')

src = Document(SRC)
sbody = src.element.body
sels = list(sbody)
assert len(sels) == 143, '整合严谨版 body len=%d 期望143' % len(sels)

cur = anchor
for se in sels[0:142]:      # 0..141，跳过142(sectPr)
    cp = copy.deepcopy(se)
    cur.addnext(cp)
    cur = cp
print('已插入整合严谨版第五部分 (元素0-141)')

doc.save(OUT)
print('已保存:', OUT)

# ---------- 校验 ----------
doc2 = Document(OUT)
els2 = list(doc2.element.body)
print('合并后 body len =', len(els2))
# 找关键标题
for i, e in enumerate(els2):
    tag = e.tag.split('}')[1]
    if tag == 'p':
        texts = [t.text or '' for t in e.iter(qn('w:t'))]
        txt = ''.join(texts).strip()
        if txt.startswith(('五、', '5.1', '5.3.7', '六、', '七、', 'AI工具', '参考文献', '附录')):
            print('  [%d] %s' % (i, txt[:40]))
