# -*- coding: utf-8 -*-
"""将 paper/paper_sections/ 的markdown章节转换为国赛格式 Word (.docx)

规范：A4、2.5cm页边距、宋体小四正文、黑体标题、三线表、嵌入图表
输出：paper/final_paper.docx
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEC_DIR = os.path.join("paper", "paper_sections")
OUT = os.path.join("paper", "final_paper.docx")
FIGURE_DIR = os.path.join("results", "figures")


def set_east_asia(run, font="宋体"):
    run.font.name = font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font)


def add_para(doc, text, size=12, font="宋体", bold=False, align=None,
             first_line_indent=None, space_after=6, line_spacing=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    # 处理加粗 **...**
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            run = p.add_run(part)
        run.font.size = Pt(size)
        set_east_asia(run, font)
        run.bold = run.bold or bold
    return p


def add_heading(doc, text, level):
    sizes = {1: 16, 2: 14, 3: 12}
    fonts = {1: "黑体", 2: "黑体", 3: "黑体"}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(8 if level == 1 else 6)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(sizes.get(level, 12))
    set_east_asia(run, fonts.get(level, "黑体"))
    return p


def add_table(doc, rows, caption=None):
    if caption:
        add_para(doc, caption, size=10.5, font="黑体", align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_after=4, line_spacing=1.2)
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    ncol = len(header)
    nrow = len(rows) - 1
    table = doc.add_table(rows=nrow + 1, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10.5)
        set_east_asia(run)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 数据行
    for i, row in enumerate(rows[1:], start=1):
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        for j in range(ncol):
            val = cells[j] if j < len(cells) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(10.5)
            set_east_asia(run)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 三线表样式：上粗线、表头下细线、下粗线
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ("top", "bottom"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "12")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for tag in ("left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)
    # 表头下细线
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "000000")
        tcBorders.append(b)
        tcPr.append(tcBorders)
    return table


def add_figure(doc, line):
    m = re.search(r"!\[(.*?)\]\((.+?)\)", line)
    if not m:
        return False
    caption, path = m.group(1), m.group(2)
    fig = os.path.join(FIGURE_DIR, os.path.basename(path.replace("\\", "/")))
    if not os.path.exists(fig):
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(fig, width=Cm(13.5))
    add_para(doc, caption, size=10.5, font="黑体", align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=10, line_spacing=1.2)
    return True


def is_table_sep(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|?\s*$", line)) and "-" in line


def convert_file(doc, fname):
    with open(os.path.join(SEC_DIR, fname), encoding="utf-8") as f:
        lines = f.read().split("\n")
    i = 0
    first_heading_done = fname.startswith("00_")
    while i < len(lines):
        ln = lines[i].strip()
        if not ln:
            i += 1
            continue
        # 表头分隔
        if ln.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = [ln]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            add_table(doc, rows)
            continue
        # 图片
        if ln.startswith("!["):
            add_figure(doc, ln)
            i += 1
            continue
        # 标题
        hm = re.match(r"^(#{1,4})\s+(.+)$", ln)
        if hm:
            level = len(hm.group(1))
            txt = re.sub(r"^[一二三四五六七八九十]+、\s*", "", hm.group(2).strip())
            txt = re.sub(r"^\d+(\.\d+)*\s*", "", txt)
            if fname.startswith("00_") and not first_heading_done:
                add_para(doc, hm.group(2).strip(), size=16, font="黑体", bold=True,
                         align=WD_ALIGN_PARAGRAPH.CENTER)
                first_heading_done = True
            else:
                add_heading(doc, txt, level)
            i += 1
            continue
        # 列表
        if re.match(r"^\s*[-*]\s+", ln):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i].strip()):
                item = re.sub(r"^\s*[-*]\s+", "", lines[i].strip())
                add_para(doc, "• " + item, first_line_indent=0.74)
                i += 1
            continue
        if re.match(r"^\s*\d+[\.、]\s+", ln):
            while i < len(lines) and re.match(r"^\s*\d+[\.、]\s+", lines[i].strip()):
                item = re.sub(r"^\s*\d+[\.、]\s+", "", lines[i].strip())
                add_para(doc, item, first_line_indent=0.74)
                i += 1
            continue
        # 普通段落（含公式，公式保留 $...$ 文本）
        para = []
        while i < len(lines) and lines[i].strip() and not (
            lines[i].strip().startswith("|") or lines[i].strip().startswith("#")
            or lines[i].strip().startswith("![") or lines[i].strip().startswith("$$")):
            para.append(lines[i].strip())
            i += 1
        if para:
            text = " ".join(para)
            add_para(doc, text, first_line_indent=0.74)
        # 公式块
        if i < len(lines) and lines[i].strip().startswith("$$"):
            if lines[i].strip().endswith("$$") and len(lines[i].strip()) > 4:
                inner = lines[i].strip()[2:-2].strip()
                add_para(doc, inner, size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                         space_after=6, line_spacing=1.3)
                i += 1
            else:
                j = i + 1
                eqs = []
                while j < len(lines) and not lines[j].strip().startswith("$$"):
                    eqs.append(lines[j].strip())
                    j += 1
                add_para(doc, " ".join(eqs), size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                         space_after=6, line_spacing=1.3)
                i = j + 1
            continue


def main():
    doc = Document()
    # 页面：A4, 2.5cm边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    # 默认样式
    style = doc.styles["Normal"]
    style.font.size = Pt(12)
    style.font.name = "宋体"
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "宋体")

    order = ["00_abstract.md", "01_problem_restatement.md", "02_problem_analysis.md",
             "03_assumptions.md", "04_symbols.md", "05_q1.md", "05_q2.md", "05_q3.md",
             "05_q4.md", "05_q5.md", "06_model_check.md", "07_evaluation.md",
             "08_references.md", "09_ai_declaration.md", "10_appendix.md"]
    for fname in order:
        if os.path.exists(os.path.join(SEC_DIR, fname)):
            convert_file(doc, fname)

    doc.save(OUT)
    print(f"final_paper.docx 已生成: {OUT} ({os.path.getsize(OUT)//1024} KB)")


if __name__ == "__main__":
    main()
