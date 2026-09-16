# -*- coding: utf-8 -*-
"""
修一处引用标点：引用应置于句号之前。
  影响。[7][8][9]。  ->  影响[7][8][9]。
"""
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

SRC = r'D:\数模工作流\delivery\C题论文_国奖附录定稿.docx'

OLD = '对运行费用的影响。[7][8][9]。'
NEW = '对运行费用的影响[7][8][9]。'

d = Document(SRC)
hit = 0
for p in d.paragraphs:
    full = ''.join(r.text for r in p.runs)
    if OLD in full:
        new_full = full.replace(OLD, NEW)
        first = p.runs[0]
        first.text = new_full
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        hit += 1
        print('修正段落:')
        print('  旧:', full[:110])
        print('  新:', new_full[:110])

# 表格内
for tbl in d.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full = ''.join(r.text for r in p.runs)
                if OLD in full:
                    first = p.runs[0]
                    first.text = full.replace(OLD, NEW)
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)
                    hit += 1

print('共修正 %d 处' % hit)
if hit:
    d.save(SRC)
    print('已保存:', SRC, os.path.getsize(SRC), 'bytes')
