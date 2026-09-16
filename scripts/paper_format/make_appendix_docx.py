# -*- coding: utf-8 -*-
"""把附录B替换为国奖样式：行号 + 语法高亮 + 无框线 + 固定行距19.3pt"""
import sys, io, os, json, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC_DOCX = r'D:\数模工作流\delivery\C题论文_附录代码终修版.docx'
OUT_DOCX = r'D:\数模工作流\tmp\C题论文_国奖附录版.docx'
DATA = json.load(open(r'D:\数模工作流\tmp\appendix_b.json', encoding='utf-8'))

# ---- 国奖实测参数 ----
LINE_TW = 386           # 固定行距 19.3pt = 386 twips
SZ = '24'               # 12pt
NUM_COLOR = '7F807F'    # 行号灰
FONT_MONO = 'Consolas'
FONT_EA = '微软雅黑'
FONT_NUM = 'Times New Roman'
COL_NUM_W = '540'       # 行号列宽 twips（27pt）。可用宽 = 540-100 = 440twips = 22pt
                        # 必须 ≥ 3位数×6pt = 18pt，否则 3 位行号折行使行高翻倍（曾致 91→103 页）
COL_CODE_W = '9000'     # 代码列宽 twips（450pt）——正文区 453.6pt 加悬挂后折行约 64.6 字符，与22C的64字符一致
COL_BLUE_W = '80'       # 右侧浅蓝带列宽 twips（4pt）—— 对齐22C：527.5->531.2pt
TBL_IND = '-400'        # 表格左移 20pt：行号挂进左边距（22C 同样做法，右缘仍留在页边距内）
BLUE_FILL = 'CED4F4'    # 22C 右侧浅蓝带实测色
RULE_SZ = '4'           # 黑线 0.5pt（22C 实测竖线宽 0.48pt）
HANG_TW = '400'         # 代码段悬挂缩进 20pt（22C 实测折行缩进 20pt）
NUM_RMAR = '100'        # 行号列右边距 5pt（列宽27pt - 5pt = 22pt 可用，容得下3位数）
CODE_LMAR = '75'        # 代码列左边距 3.75pt（22C 实测代码起 71.3 vs 分隔线 67.55）
BORDER_COLOR = '7F807F'

_order = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
          'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
          'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE',
          'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
          'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection',
          'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
          'sectPr', 'pPrChange']
_tc_order = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd',
             'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']


def ins_ordered(parent, child, order):
    tag = child.tag.split('}')[-1]
    try:
        idx = order.index(tag)
    except ValueError:
        parent.append(child)
        return child
    for existing in list(parent):
        etag = existing.tag.split('}')[-1]
        try:
            eidx = order.index(etag)
        except ValueError:
            continue
        if eidx > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def set_spacing(p, line=LINE_TW, before='0', after='0'):
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = ins_ordered(pPr, OxmlElement('w:spacing'), _order)
    sp.set(qn('w:line'), str(line))
    sp.set(qn('w:lineRule'), 'exact')
    sp.set(qn('w:before'), before)
    sp.set(qn('w:after'), after)
    return sp


def set_run_font(run, color, font_mid=FONT_MONO, size=SZ, num=False):
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = ins_ordered(rPr, OxmlElement('w:rFonts'), ['rStyle', 'rFonts', 'b', 'bCs', 'i',
                                                        'iCs', 'caps', 'smallCaps', 'strike',
                                                        'dstrike', 'outline', 'shadow',
                                                        'emboss', 'imprint', 'noProof',
                                                        'snapToGrid', 'vanish', 'webHidden',
                                                        'color', 'spacing', 'w', 'kern',
                                                        'position', 'sz', 'szCs', 'highlight',
                                                        'u', 'effect', 'bdr', 'shd',
                                                        'fitText', 'vertAlign', 'rtl', 'cs',
                                                        'em', 'lang', 'eastAsianLayout',
                                                        'specVanish', 'oMath'])
    if num:
        rf.set(qn('w:ascii'), FONT_NUM)
        rf.set(qn('w:hAnsi'), FONT_NUM)
        rf.set(qn('w:eastAsia'), FONT_NUM)
    else:
        rf.set(qn('w:ascii'), FONT_MONO)
        rf.set(qn('w:hAnsi'), FONT_MONO)
        rf.set(qn('w:eastAsia'), FONT_EA)
    rf.set(qn('w:cs'), FONT_MONO)
    c = rPr.find(qn('w:color'))
    if c is None:
        c = ins_ordered(rPr, OxmlElement('w:color'), ['rStyle', 'rFonts', 'b', 'bCs', 'i',
                                                      'iCs', 'caps', 'smallCaps', 'strike',
                                                      'dstrike', 'outline', 'shadow',
                                                      'emboss', 'imprint', 'noProof',
                                                      'snapToGrid', 'vanish', 'webHidden',
                                                      'color', 'spacing', 'w', 'kern',
                                                      'position', 'sz', 'szCs', 'highlight',
                                                      'u', 'effect', 'bdr', 'shd', 'fitText',
                                                      'vertAlign', 'rtl', 'cs', 'em', 'lang',
                                                      'eastAsianLayout', 'specVanish',
                                                      'oMath'])
    c.set(qn('w:val'), color)
    for tag in ('w:sz', 'w:szCs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = OxmlElement(tag)
            rPr.append(e)
        e.set(qn('w:val'), str(size))


def set_no_snap(p):
    """关闭网格吸附：docGrid linePitch=360(18pt) 会把 19.3pt 行距撑成 2 格=36pt"""
    pPr = p._p.get_or_add_pPr()
    sg = OxmlElement('w:snapToGrid')
    sg.set(qn('w:val'), '0')
    ins_ordered(pPr, sg, _order)


def set_ind_zero(p, left='0', hanging=None):
    """清掉继承自正文样式的首行缩进 480twips(24pt)。
    hanging 不为 None 时，做成悬挂缩进（首行在 left-hanging，折行在 left）——22C 折行缩进 20pt。"""
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:ind'))
    if old is not None:
        pPr.remove(old)
    ind = OxmlElement('w:ind')
    if hanging:
        ind.set(qn('w:left'), str(left))
        ind.set(qn('w:hanging'), str(hanging))
    else:
        ind.set(qn('w:firstLine'), '0')
        ind.set(qn('w:left'), str(left))
        ind.set(qn('w:right'), '0')
    ins_ordered(pPr, ind, _order)


def cell_text(cell, spans, align, num=False, size=SZ, snap_off=True, hang=False):
    """spans: [(text, color)]；返回段落"""
    p = cell.paragraphs[0]
    if snap_off:
        set_no_snap(p)
    if hang:
        set_ind_zero(p, left=HANG_TW, hanging=HANG_TW)
    else:
        set_ind_zero(p)
    pPr = p._p.get_or_add_pPr()
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), align)
    ins_ordered(pPr, jc, _order)
    set_spacing(p)
    for txt, col in spans:
        if txt == '':
            continue
        r = p.add_run(txt)
        set_run_font(r, col, num=num, size=size)
    return p


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """按需设置单元格边框。spec = (val, sz, color) 或 ('nil', None, None)"""
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:tcBorders'))
    if old is not None:
        tcPr.remove(old)
    tcB = OxmlElement('w:tcBorders')
    for side, spec in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        if spec is None:
            continue
        val, sz, color = spec
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), val)
        if val != 'nil':
            e.set(qn('w:sz'), sz)
            e.set(qn('w:space'), '0')
            e.set(qn('w:color'), color)
        tcB.append(e)
    ins_ordered(tcPr, tcB, _tc_order)


def set_cell_borders_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcB = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'nil')
        tcB.append(e)
    ins_ordered(tcPr, tcB, _tc_order)


def set_cell_margins(cell, left='0', right='60'):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, val in (('top', '0'), ('left', left), ('bottom', '0'), ('right', right)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), val)
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    ins_ordered(tcPr, mar, _tc_order)


def set_valign(cell, val='center'):
    tcPr = cell._tc.get_or_add_tcPr()
    e = OxmlElement('w:vAlign')
    e.set(qn('w:val'), val)
    ins_ordered(tcPr, e, _tc_order)


def set_cell_shd(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    ins_ordered(tcPr, shd, _tc_order)


def set_tbl_gj_borders(tbl):
    """国奖附录B表格线：只保留 insideV 黑色分隔线(行号列|代码列|右侧带)。
    上下横线不能放在表格级：Word 会在每个跨页分片的页首页尾重复画，22C 的横线其实来自文件名段落的上下边框。"""
    tblPr = tbl._tbl.tblPr
    for tag in ('w:tblBorders', 'w:tblInd', 'w:tblW'):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    spec = [('top', 'nil'), ('left', 'nil'), ('bottom', 'nil'), ('right', 'nil'),
            ('insideH', 'nil'), ('insideV', 'single')]
    for side, val in spec:
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), val)
        if val == 'single':
            e.set(qn('w:sz'), RULE_SZ)
            e.set(qn('w:space'), '0')
            e.set(qn('w:color'), '000000')
        b.append(e)
    tblPr.append(b)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(int(COL_NUM_W) + int(COL_CODE_W) + int(COL_BLUE_W)))
    tw.set(qn('w:type'), 'dxa')
    tblPr.append(tw)
    ind = OxmlElement('w:tblInd')
    ind.set(qn('w:w'), TBL_IND)
    ind.set(qn('w:type'), 'dxa')
    tblPr.append(ind)
    lay = tblPr.find(qn('w:tblLayout'))
    if lay is None:
        lay = OxmlElement('w:tblLayout')
        tblPr.append(lay)
    lay.set(qn('w:type'), 'fixed')


def add_hash_border(p, top=False, bottom=False, sz=RULE_SZ):
    """给段落加国奖style上下横线（22C 实测：文件名行上下各一条黑线，段高约 30pt）"""
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for side, want in (('top', top), ('bottom', bottom)):
        if not want:
            continue
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '1')
        e.set(qn('w:color'), '000000')
        pbdr.append(e)
    ins_ordered(pPr, pbdr, _order)


def set_tr_props(row):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement('w:cantSplit')
    trPr.append(cs)


def add_bottom_border(p, sz='6'):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single')
    bt.set(qn('w:sz'), sz)
    bt.set(qn('w:space'), '1')
    bt.set(qn('w:color'), BORDER_COLOR)
    pbdr.append(bt)
    ins_ordered(pPr, pbdr, _order)


def main():
    # 自检：行号列可用宽度必须容得下最长的行号，否则行号折行 → 行高翻倍 → 页数暴涨
    maxno = max((len(str(no)) for it in DATA for no, _, _ in it['lines'] if no is not None), default=1)
    usable_tw = int(COL_NUM_W) - int(NUM_RMAR)
    need_tw = maxno * 120          # TNR 12pt 数字宽 0.5em = 6pt = 120 twips
    print('最长行号 %d 位，需 %d twips，行号列可用 %d twips' % (maxno, need_tw, usable_tw))
    assert usable_tw >= need_tw + 40, '行号列太窄：%d 位行号会折行' % maxno

    shutil.copyfile(SRC_DOCX, OUT_DOCX)
    doc = Document(OUT_DOCX)
    body = doc.element.body

    # 1) 定位锚点「附录 B　源程序代码」
    anchor = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('附录 B') or t.startswith('附录B'):
            anchor = p._p
            break
    if anchor is None:
        raise SystemExit('未找到「附录 B 源程序代码」标题段落')
    print('锚点段落:', [p.text for p in doc.paragraphs if p._p is anchor][0][:40])

    # 2) 删除锚点之后的全部内容（保留 sectPr）
    elems = list(body)
    i = elems.index(anchor)
    removed = 0
    for e in elems[i + 1:]:
        if e.tag == qn('w:sectPr'):
            continue
        body.remove(e)
        removed += 1
    print('删除旧内容元素', removed, '个')

    # 3) 依次插入新内容
    cursor = anchor
    n_files = len(DATA)
    for fi, it in enumerate(DATA):
        # 代码表格（3列：行号 | 代码 | 右侧浅蓝带）
        tbl = doc.add_table(rows=0, cols=3)
        cursor.addnext(tbl._tbl)
        cursor = tbl._tbl
        set_tbl_gj_borders(tbl)
        grid = tbl._tbl.find(qn('w:tblGrid'))
        if grid is not None:
            cols = grid.findall(qn('w:gridCol'))
            for i, w in enumerate((COL_NUM_W, COL_CODE_W, COL_BLUE_W)):
                if i < len(cols):
                    cols[i].set(qn('w:w'), w)

        # 3a) 文件名行：3列合并成一行，上下各一条黑横线（横线宽=表宽，与22C一致）
        hrow = tbl.add_row()
        set_tr_props(hrow)
        hcell = hrow.cells[0].merge(hrow.cells[2])
        set_cell_borders(hcell, top=('single', RULE_SZ, '000000'),
                         bottom=('single', RULE_SZ, '000000'))
        set_valign(hcell, 'center')
        set_cell_margins(hcell, left='0', right='0')
        hp = hcell.paragraphs[0]
        set_no_snap(hp)
        set_ind_zero(hp)
        set_spacing(hp, line='280', before='160', after='160')
        r = hp.add_run(it['name'])
        set_run_font(r, '000000', num=False, size=SZ)
        r.bold = True
        r2 = hp.add_run('　' + it['title'])
        set_run_font(r2, '404040', num=False, size='21')

        for no, text, spans in it['lines']:
            row = tbl.add_row()
            set_tr_props(row)
            c0, c1, c2 = row.cells
            # 不设单元格级边框：否则会覆盖表格级的 insideV 分隔线
            set_valign(c0, 'top')
            set_valign(c1, 'top')
            set_valign(c2, 'top')
            set_cell_margins(c0, left='0', right=NUM_RMAR)
            set_cell_margins(c1, left=CODE_LMAR, right='0')
            set_cell_margins(c2, left='0', right='0')
            set_cell_shd(c2, BLUE_FILL)
            # 列宽
            for c, w in ((c0, COL_NUM_W), (c1, COL_CODE_W), (c2, COL_BLUE_W)):
                tcPr = c._tc.get_or_add_tcPr()
                e = tcPr.find(qn('w:tcW'))
                if e is None:
                    e = OxmlElement('w:tcW')
                    ins_ordered(tcPr, e, _tc_order)
                e.set(qn('w:w'), w)
                e.set(qn('w:type'), 'dxa')
            cell_text(c0, [(('' if no is None else str(no)), NUM_COLOR)], 'right', num=True)
            cell_text(c1, spans, 'left', hang=True)
            # 右侧浅蓝带：压扁段落，避免撑高行
            pz = c2.paragraphs[0]
            set_no_snap(pz)
            set_ind_zero(pz)
            set_spacing(pz, line='20', before='0', after='0')

        # 3b) 最后一个文件的末行补下边线，形成收边（22C 末页同样有一条收边线）
        if fi == n_files - 1:
            last = tbl.rows[-1]
            for c in last.cells:
                _tcPr = c._tc.get_or_add_tcPr()
                _b = _tcPr.find(qn('w:tcBorders'))
                if _b is None:
                    _b = OxmlElement('w:tcBorders')
                    ins_ordered(_tcPr, _b, _tc_order)
                e = OxmlElement('w:bottom')
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), RULE_SZ)
                e.set(qn('w:space'), '0')
                e.set(qn('w:color'), '000000')
                _b.append(e)

    doc.save(OUT_DOCX)
    print('已保存:', OUT_DOCX, os.path.getsize(OUT_DOCX), 'bytes')


main()
