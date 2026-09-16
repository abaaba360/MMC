from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"D:\数模工作流\delivery\C题论文.docx")
OUTPUT = Path(r"D:\数模工作流\delivery\C题论文_支撑材料表修订版.docx")

ROWS = [
    ("Q1.py", "问题一求解程序"),
    ("Q2.py", "问题二求解程序"),
    ("Q3.py", "问题三求解程序"),
    ("Q4.py", "问题四求解程序"),
    ("common.py", "公共模型与求解函数"),
    ("data_q1.py", "问题一数据读取程序"),
    ("data.py", "公共数据处理程序"),
    ("forecast.py", "光伏与负荷预测程序"),
    ("rolling.py", "滚动优化调度程序"),
    ("check.py", "四问结果校验程序"),
    ("export.py", "结果工作簿导出程序"),
    ("export_matlab.py", "MATLAB 绘图数据导出程序"),
    ("plot_q1.m", "问题一结果绘图程序"),
    ("plot_figures.m", "问题二至四绘图程序"),
    ("run_all.py", "全部程序运行入口"),
    ("requirements.txt", "Python 依赖环境清单"),
    ("results/", "计算结果与校验文件"),
    ("figures/", "论文使用的图件"),
    ("代码运行说明.txt", "程序运行方法说明"),
    ("AI工具使用详情.pdf", "AI 工具使用记录"),
    ("文件清单.json", "支撑材料文件清单"),
]


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, settings in edges.items():
        tag = qn(f"w:{edge_name}")
        edge = tc_borders.find(tag)
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key, value in settings.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for name, val, size in (
        ("top", "single", "12"),
        ("left", "nil", "0"),
        ("bottom", "single", "12"),
        ("right", "nil", "0"),
        ("insideH", "nil", "0"),
        ("insideV", "nil", "0"),
    ):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), val)
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), "000000")
        borders.append(edge)
    tbl_pr.append(borders)


def clear_cell_shading(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is not None:
        tc_pr.remove(shd)


def set_cell_margins(cell, top=45, start=95, bottom=45, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def no_wrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:noWrap"))
    if node is None:
        tc_pr.append(OxmlElement("w:noWrap"))


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), "宋体")


def set_cell_text(cell, text, *, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    clear_cell_shading(cell)
    set_cell_margins(cell)
    no_wrap(cell)


def find_target_table(document):
    for table in document.tables:
        joined = " ".join(cell.text for row in table.rows for cell in row.cells)
        if "Q1.py" in joined and "AI工具使用详情" in joined:
            return table
    raise RuntimeError("未找到 A.1 支撑材料文件列表")


def main():
    document = Document(SOURCE)
    table = find_target_table(document)

    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(ROWS) + 1:
        table.add_row()

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)

    widths = (Cm(7.3), Cm(8.0))
    headers = ("文件名", "功能描述")
    for col_index, text in enumerate(headers):
        cell = table.cell(0, col_index)
        cell.width = widths[col_index]
        set_cell_text(cell, text, bold=False)
        set_cell_border(cell, bottom={"val": "single", "sz": "8", "space": "0", "color": "000000"})

    table.rows[0].height = Cm(0.68)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    set_repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_index, values in enumerate(ROWS, start=1):
        row = table.rows[row_index]
        row.height = Cm(0.55)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = widths[col_index]
            set_cell_text(cell, value)

    # Remove any cell-level side or internal borders left by the former grid table.
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            edges = {
                "left": {"val": "nil", "sz": "0", "space": "0", "color": "000000"},
                "right": {"val": "nil", "sz": "0", "space": "0", "color": "000000"},
                "top": {"val": "nil", "sz": "0", "space": "0", "color": "000000"},
                "bottom": {"val": "nil", "sz": "0", "space": "0", "color": "000000"},
            }
            if row_index == 0:
                edges["top"] = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
                edges["bottom"] = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
            if row_index == len(table.rows) - 1:
                edges["bottom"] = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
            set_cell_border(cell, **edges)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
